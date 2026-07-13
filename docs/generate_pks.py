"""
Generate template Perjanjian Kerja Sama (PKS) untuk implementasi & penggunaan
aplikasi TMS AR Cargo Business Suite dalam format MS Word (.docx).

Cara pakai:
    python3 generate_pks.py

Output: /app/docs/PKS_Aplikasi_AR_Cargo_TMS.docx
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
            tcBorders.append(e)
    tc_pr.append(tcBorders)


def add_heading(doc, text, level=1, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
    elif level == 2:
        run.font.size = Pt(12)
    else:
        run.font.size = Pt(11)
    return p


def add_para(doc, text, bold=False, italic=False, indent=0, align=None, size=11):
    p = doc.add_paragraph()
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_pasal(doc, no, judul):
    """Header pasal: PASAL X - JUDUL, center-bold."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"PASAL {no}\n")
    r1.bold = True
    r1.font.size = Pt(12)
    r2 = p.add_run(judul.upper())
    r2.bold = True
    r2.underline = True
    r2.font.size = Pt(12)


def add_ayat_list(doc, items, start=1):
    """Ayat bernomor (1) (2) (3) — text justify."""
    for i, item in enumerate(items, start=start):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        run = p.add_run(f"({i})  {item}")
        run.font.size = Pt(11)


def add_bullet_list(doc, items, indent=1.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(indent)
        run = p.runs[0] if p.runs else p.add_run("")
        p.runs[0].text = it if p.runs else it
        for r in p.runs:
            r.font.size = Pt(11)


def build():
    doc = Document()

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Margins
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # ====================== JUDUL ======================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PERJANJIAN KERJA SAMA")
    r.bold = True; r.font.size = Pt(16); r.underline = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PENYEDIAAN & PENGGUNAAN APLIKASI\n"
                  "TRANSPORT MANAGEMENT SYSTEM (TMS) — AR CARGO BUSINESS SUITE")
    r.bold = True; r.font.size = Pt(13)

    add_para(doc, "Nomor: ____/PKS/______/______", align="center", bold=True)
    add_para(doc, "")

    # ====================== PEMBUKAAN ======================
    intro = ("Perjanjian Kerja Sama ini (selanjutnya disebut “PERJANJIAN”) dibuat "
             "dan ditandatangani pada hari ______, tanggal ____ bulan ______________ "
             "tahun __________ (___-___-____), oleh dan antara:")
    add_para(doc, intro, align="justify")

    # PIHAK PERTAMA
    add_para(doc, "")
    add_para(doc, "PIHAK PERTAMA", bold=True)
    pihak1 = [
        "Nama Perusahaan\t:\t_______________________________________________",
        "Alamat\t\t:\t_______________________________________________",
        "Diwakili oleh\t:\t_______________________________________________",
        "Jabatan\t\t:\t_______________________________________________",
    ]
    for x in pihak1:
        add_para(doc, x, indent=0.5)
    add_para(doc,
             "Dalam hal ini bertindak untuk dan atas nama perusahaan tersebut di atas selaku "
             "PENYEDIA APLIKASI, selanjutnya disebut sebagai “PIHAK PERTAMA”.",
             align="justify")

    # PIHAK KEDUA
    add_para(doc, "")
    add_para(doc, "PIHAK KEDUA", bold=True)
    pihak2 = [
        "Nama Perusahaan\t:\tPT. Abirama Mitra Pratama / PT. Abirama Mitra Pradana",
        "Alamat\t\t:\tGrand Galaxy City RSK 6 No. 73 Jakasetia, Bekasi Selatan,",
        "\t\t\tKota Bekasi, Jawa Barat 17147",
        "Diwakili oleh\t:\t_______________________________________________",
        "Jabatan\t\t:\t_______________________________________________",
    ]
    for x in pihak2:
        add_para(doc, x, indent=0.5)
    add_para(doc,
             "Dalam hal ini bertindak untuk dan atas nama perusahaan tersebut di atas selaku "
             "PENGGUNA APLIKASI, selanjutnya disebut sebagai “PIHAK KEDUA”.",
             align="justify")

    add_para(doc, "")
    add_para(doc,
             "PIHAK PERTAMA dan PIHAK KEDUA secara bersama-sama selanjutnya disebut sebagai "
             "“PARA PIHAK”, dan masing-masing sebagai “PIHAK”.",
             align="justify")

    # ====================== MENIMBANG ======================
    add_para(doc, "")
    add_heading(doc, "MENIMBANG:", level=2)
    menimbang = [
        "Bahwa PIHAK PERTAMA adalah pihak yang mengembangkan, memelihara, dan menyediakan "
        "aplikasi berbasis web bernama “AR Cargo Business Suite” yang selanjutnya disebut “APLIKASI”, "
        "yang berfungsi sebagai sistem manajemen operasional transportasi & keuangan (Transport "
        "Management System / TMS) beserta modul-modul pendukungnya;",

        "Bahwa PIHAK KEDUA adalah perusahaan yang bergerak di bidang jasa pengiriman "
        "barang (freight forwarder / cargo) dan membutuhkan APLIKASI untuk mendukung kegiatan "
        "operasional, keuangan, akuntansi, dan pelaporan harian;",

        "Bahwa berdasarkan hal-hal tersebut di atas, PARA PIHAK sepakat mengadakan Perjanjian "
        "Kerja Sama dengan syarat dan ketentuan sebagaimana diatur dalam pasal-pasal berikut.",
    ]
    add_ayat_list(doc, menimbang)

    # ====================== PASAL 1 – DEFINISI ======================
    doc.add_page_break()
    add_pasal(doc, 1, "Definisi")
    definisi = [
        "“APLIKASI” adalah perangkat lunak berbasis web AR Cargo Business Suite (TMS) yang disediakan "
        "oleh PIHAK PERTAMA dan diakses PIHAK KEDUA melalui browser atau URL yang ditentukan.",

        "“DATA” adalah seluruh data operasional, transaksi, dokumen, dan informasi milik PIHAK KEDUA "
        "yang di-input, di-upload, atau dihasilkan oleh APLIKASI.",

        "“PENGGUNA” adalah karyawan, staf, atau pihak yang diberi otorisasi oleh PIHAK KEDUA "
        "untuk mengakses APLIKASI menggunakan kredensial (username & password).",

        "“LAYANAN” adalah keseluruhan fungsi, fitur, hosting, pemeliharaan, dan dukungan teknis "
        "yang disediakan PIHAK PERTAMA sehubungan dengan APLIKASI.",

        "“MASA BERLAKU” adalah jangka waktu di mana PIHAK KEDUA berhak menggunakan APLIKASI "
        "sebagaimana diatur pada Pasal 9 Perjanjian ini.",

        "“HARI KERJA” adalah hari Senin sampai dengan Jumat, tidak termasuk hari libur nasional "
        "di Indonesia, pukul 09.00 – 17.00 WIB.",
    ]
    add_ayat_list(doc, definisi)

    # ====================== PASAL 2 – MAKSUD & TUJUAN ======================
    add_pasal(doc, 2, "Maksud dan Tujuan")
    add_ayat_list(doc, [
        "Maksud Perjanjian ini adalah untuk mengatur hak dan kewajiban PARA PIHAK sehubungan "
        "dengan penyediaan APLIKASI oleh PIHAK PERTAMA kepada PIHAK KEDUA.",

        "Tujuan Perjanjian ini adalah agar PIHAK KEDUA dapat menggunakan APLIKASI untuk mendukung "
        "kegiatan operasional dan administrasi bisnis pengiriman barang (cargo), termasuk namun "
        "tidak terbatas pada: pengelolaan Airway Bill (AWB), invoice, pembayaran, biaya operasional, "
        "hutang vendor, akuntansi dasar (PSAK), pelaporan laba-rugi, dan modul HR/Payroll.",
    ])

    # ====================== PASAL 3 – RUANG LINGKUP APLIKASI ======================
    add_pasal(doc, 3, "Ruang Lingkup Aplikasi")
    add_para(doc,
             "PIHAK PERTAMA menyediakan APLIKASI dengan modul-modul utama sebagai berikut:",
             align="justify")

    modul_data = [
        ("A. ORDER TO CASH (O2C)", [
            "Input Order — pencatatan Airway Bill (AWB) baru meliputi identitas customer, "
            "rute (origin-destination), service (LCL/FCL/LTL/FTL/AIR/Lumpsum), tipe kontainer "
            "(20FT/40FT/HQ/Open Top), armada, jadwal keberangkatan (ETD/ETA), tarif & pajak (PPN 0/1.1/11%);",

            "Daftar AWB — daftar seluruh order dengan filter status, customer, periode; "
            "tracking posisi barang real-time; validasi anti-duplikasi No. AWB; "
            "penerimaan barang (POD);",

            "Rekap & Rekonsiliasi — verifikasi biaya vs tarif per AWB, penentuan status siap-invoice;",

            "Buat Invoice — pembuatan invoice tunggal maupun bulk multi-AWB, kop surat "
            "template dapat dikustomisasi, cetak/preview HTML & export Excel;",

            "Invoice Tercetak — arsip invoice, cetak ulang, pembatalan;",

            "Pembayaran — pencatatan pelunasan invoice, auto-jurnal kas/bank, "
            "tracking status Open/Partial/Paid.",
        ]),
        ("B. BIAYA & PENGADAAN", [
            "Plotting Biaya — assign biaya vendor per AWB, split biaya, kalkulasi HPP;",

            "Settle Hutang Vendor — pelunasan hutang ke vendor, tracking outstanding;",

            "Panjar Kerja — advance / cash advance untuk kegiatan operasional beserta "
            "settlement dan return;",

            "PO / Surat Perintah Vendor — pembuatan Purchase Order ke vendor pengangkut;",

            "Expense — pencatatan biaya umum non-vendor beserta kategori;",

            "Cost Automatic (Recurring) — biaya berulang otomatis di-post ke jurnal.",
        ]),
        ("C. MASTER & LAPORAN", [
            "Dashboard — ringkasan operasional (AWB Open, Rekonsiliasi, Invoiced, Lunas, "
            "Total Pendapatan, Total Biaya, Hutang Vendor);",

            "Dashboard Analitik — filter periode, customer, service, region; grafik trend & detail AWB;",

            "Customer — master data pelanggan, kontrak, NPWP, alamat;",

            "Perjanjian Rate — kontrak tarif per customer (mengunci tarif otomatis saat Input Order);",

            "Master Kategori Expense — kategori biaya untuk keperluan reporting & jurnal;",

            "Laba/Rugi (Order) — analisa profitabilitas per AWB.",
        ]),
        ("D. AKUNTANSI (PSAK)", [
            "Bagan Akun (Chart of Accounts) — sesuai standar PSAK Indonesia;",

            "Jurnal Umum — pencatatan manual & otomatis dari transaksi (AR, AP, Kas, Bank);",

            "Buku Besar (General Ledger) — saldo running per akun;",

            "Neraca Percobaan (Trial Balance);",

            "Laporan Keuangan — Laba/Rugi & Neraca dengan filter periode;",

            "Rekonsiliasi Bank — matching transaksi kas/bank dgn statement.",
        ]),
        ("E. HR & PAYROLL", [
            "Master Karyawan — data lengkap: identitas, kontak, legal (KTP/NPWP/KK), "
            "logistik (kendaraan, SIM), finansial (bank), kontak darurat, data keluarga, kepegawaian;",

            "Payroll — perhitungan gaji bulanan, auto-jurnal beban gaji (6101) dan PPh21 (2220);",

            "Slip Gaji — cetak per karyawan.",
        ]),
        ("F. UMUM & KEAMANAN", [
            "Multi-User dengan Role-Based Access Control (Admin, Operation, Finance, HR, Viewer, dsb);",

            "Feature-Flag per modul (Admin dapat mengaktifkan/menonaktifkan modul);",

            "Audit Trail (created_at/updated_at, created_by);",

            "Export Excel per modul;",

            "Settings — kop surat invoice, tarif pajak, logo, template dokumen.",
        ]),
    ]
    for judul, items in modul_data:
        add_para(doc, "")
        add_para(doc, judul, bold=True, size=11)
        for it in items:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(1.5)
            for r in p.runs: r.text = ""
            run = p.add_run(it) if not p.runs else p.runs[0]
            run.text = it
            run.font.size = Pt(11)

    # ====================== PASAL 4 – BATASAN & LIMITASI ======================
    doc.add_page_break()
    add_pasal(doc, 4, "Batasan dan Limitasi")
    add_para(doc,
             "PARA PIHAK sepakat bahwa penyediaan APLIKASI tunduk pada batasan-batasan berikut:",
             align="justify")
    limitasi = [
        "APLIKASI disediakan “AS IS” (sebagaimana adanya) sesuai fitur yang telah ada pada saat "
        "penandatanganan Perjanjian ini. Pengembangan fitur baru di luar cakupan akan dibicarakan "
        "secara terpisah dan dapat dikenakan biaya tambahan.",

        "APLIKASI dirancang khusus untuk kebutuhan operasional freight forwarder / cargo AR Cargo "
        "dan tidak dijamin cocok untuk industri atau model bisnis lain di luar itu.",

        "APLIKASI membutuhkan koneksi internet stabil dan browser modern (Chrome / Edge / Safari "
        "versi terbaru). PIHAK PERTAMA tidak bertanggung jawab atas gangguan yang disebabkan oleh "
        "infrastruktur jaringan atau perangkat PIHAK KEDUA.",

        "Uptime target APLIKASI adalah 99% per bulan (di luar jadwal pemeliharaan terjadwal). "
        "Pemeliharaan terjadwal akan diberitahukan minimal 24 jam sebelumnya. "
        "Downtime akibat pihak ketiga (hosting provider, DNS, ISP) berada di luar tanggung jawab PIHAK PERTAMA.",

        "APLIKASI mendukung penyimpanan data hingga kapasitas yang wajar untuk operasional harian "
        "AR Cargo. Kebutuhan penyimpanan atau bandwidth di luar batas normal (misalnya arsip video, "
        "file besar > 10MB per upload) dapat dikenakan biaya tambahan.",

        "APLIKASI TIDAK termasuk: integrasi realtime dengan sistem eksternal (marketplace, "
        "aplikasi armada pihak ketiga, EDI); modul tax e-Faktur / e-Bupot resmi DJP; sertifikasi "
        "ISO / audit keuangan formal. Setiap kebutuhan integrasi baru harus disepakati terpisah.",

        "Laporan keuangan yang dihasilkan APLIKASI merupakan alat bantu manajerial dan bukan "
        "pengganti laporan keuangan resmi yang harus diaudit oleh Kantor Akuntan Publik (KAP) "
        "sesuai peraturan perundang-undangan.",

        "PIHAK PERTAMA tidak bertanggung jawab atas kesalahan input data, kesalahan setting tarif, "
        "kesalahan konfigurasi akun, atau interpretasi laporan oleh PENGGUNA. Data yang dihasilkan "
        "APLIKASI hanya sebaik data yang di-input.",

        "APLIKASI dapat mengalami perubahan (update, patch) dari waktu ke waktu untuk perbaikan "
        "bug, keamanan, atau penambahan fitur. Update tidak boleh mengurangi fungsi yang telah "
        "disepakati tanpa kesepakatan tertulis.",

        "Kapasitas pengguna simultan (concurrent users) dan batas transaksi bulanan akan diatur "
        "dalam Lampiran teknis. Kelebihan kuota dapat dikenakan biaya tambahan.",
    ]
    add_ayat_list(doc, limitasi)

    # ====================== PASAL 5 – HAK & KEWAJIBAN ======================
    add_pasal(doc, 5, "Hak dan Kewajiban Para Pihak")

    add_para(doc, "5.1  Hak dan Kewajiban PIHAK PERTAMA:", bold=True)
    add_ayat_list(doc, [
        "Menyediakan APLIKASI beserta dokumentasi penggunaan (user manual) kepada PIHAK KEDUA.",
        "Melakukan pemeliharaan rutin, monitoring, dan backup data harian.",
        "Memberikan dukungan teknis (technical support) via email/chat pada Hari Kerja "
        "dengan response time maksimal 1x24 jam.",
        "Menjaga kerahasiaan Data PIHAK KEDUA sebagaimana diatur pada Pasal 8.",
        "Berhak menerima pembayaran biaya Layanan sesuai kesepakatan.",
        "Berhak menangguhkan Layanan bila PIHAK KEDUA lalai membayar sebagaimana Pasal 6.",
    ])

    add_para(doc, "5.2  Hak dan Kewajiban PIHAK KEDUA:", bold=True)
    add_ayat_list(doc, [
        "Menggunakan APLIKASI hanya untuk kegiatan bisnis internal PIHAK KEDUA sesuai tujuan Perjanjian.",
        "Menyediakan dan mengelola akun PENGGUNA (username & password) secara bertanggung jawab. "
        "Setiap aktivitas yang dilakukan menggunakan kredensial PENGGUNA adalah tanggung jawab PIHAK KEDUA.",
        "Menjaga akurasi & kelengkapan Data yang di-input ke APLIKASI.",
        "Membayar biaya Layanan tepat waktu sebagaimana Pasal 6.",
        "Tidak melakukan reverse engineering, decompile, atau menyalin kode sumber APLIKASI.",
        "Tidak mendistribusikan, menjual kembali, atau memberikan akses APLIKASI kepada pihak ketiga "
        "di luar afiliasi grup usaha PIHAK KEDUA tanpa persetujuan tertulis PIHAK PERTAMA.",
        "Bertanggung jawab atas kepatuhan penggunaan APLIKASI terhadap peraturan perundang-undangan yang berlaku.",
    ])

    # ====================== PASAL 6 – BIAYA & PEMBAYARAN ======================
    add_pasal(doc, 6, "Biaya dan Pembayaran")
    add_ayat_list(doc, [
        "Biaya Layanan berlaku sesuai skema yang disepakati (Setup Fee, Subscription Fee bulanan/tahunan, "
        "atau paket lainnya) sebagaimana diatur dalam Lampiran Perjanjian ini.",

        "Pembayaran dilakukan melalui transfer ke rekening yang ditunjuk PIHAK PERTAMA, paling lambat "
        "tanggal 10 pada bulan berjalan (untuk subscription bulanan) atau sesuai invoice yang diterbitkan.",

        "Keterlambatan pembayaran lebih dari 14 (empat belas) hari kalender dari tanggal jatuh tempo "
        "dapat mengakibatkan penangguhan sementara Layanan sampai kewajiban dilunasi.",

        "Semua biaya sudah termasuk PPN sesuai peraturan yang berlaku, kecuali dinyatakan lain.",

        "Biaya tambahan (custom development, integrasi baru, training on-site) akan dihitung terpisah "
        "berdasarkan Surat Penawaran Harga yang harus disetujui PARA PIHAK sebelum pekerjaan dimulai.",
    ])

    # ====================== PASAL 7 – DATA & KEPEMILIKAN ======================
    add_pasal(doc, 7, "Data dan Kepemilikan Kekayaan Intelektual")
    add_ayat_list(doc, [
        "Seluruh DATA yang di-input dan dihasilkan oleh PIHAK KEDUA di dalam APLIKASI adalah milik "
        "sepenuhnya PIHAK KEDUA.",

        "Hak Kekayaan Intelektual (HAKI) atas APLIKASI, termasuk kode sumber, desain, arsitektur, "
        "dan seluruh dokumentasi teknis, adalah milik sepenuhnya PIHAK PERTAMA.",

        "PIHAK PERTAMA memberikan lisensi non-eksklusif, tidak dapat dialihkan, kepada PIHAK KEDUA "
        "untuk menggunakan APLIKASI selama Masa Berlaku Perjanjian.",

        "PIHAK KEDUA berhak melakukan export DATA sewaktu-waktu (dalam format Excel/CSV) melalui "
        "fitur export yang tersedia.",

        "Setelah berakhirnya Perjanjian, PIHAK PERTAMA wajib menyediakan salinan DATA PIHAK KEDUA "
        "dalam format Excel/CSV dan menghapus DATA dari server dalam waktu 30 (tiga puluh) hari.",
    ])

    # ====================== PASAL 8 – KERAHASIAAN ======================
    doc.add_page_break()
    add_pasal(doc, 8, "Kerahasiaan")
    add_ayat_list(doc, [
        "PARA PIHAK sepakat menjaga kerahasiaan seluruh informasi rahasia yang diperoleh berkaitan "
        "dengan pelaksanaan Perjanjian ini.",

        "Informasi rahasia mencakup namun tidak terbatas pada: DATA operasional & keuangan PIHAK KEDUA, "
        "kode sumber & arsitektur APLIKASI, strategi bisnis, harga khusus, dan rencana kerja.",

        "Kewajiban kerahasiaan tetap berlaku selama 3 (tiga) tahun sejak berakhirnya Perjanjian.",

        "Pelanggaran kewajiban kerahasiaan dapat dikenakan ganti rugi sesuai kerugian nyata yang timbul.",
    ])

    # ====================== PASAL 9 – JANGKA WAKTU ======================
    add_pasal(doc, 9, "Jangka Waktu")
    add_ayat_list(doc, [
        "Perjanjian ini berlaku selama ____ (____________) tahun terhitung sejak tanggal "
        "penandatanganan dan dapat diperpanjang secara otomatis kecuali salah satu PIHAK "
        "memberitahukan keinginan untuk tidak memperpanjang minimal 30 (tiga puluh) hari sebelum berakhirnya Masa Berlaku.",

        "Perpanjangan tunduk pada kesepakatan biaya Layanan yang baru bila terdapat penyesuaian.",
    ])

    # ====================== PASAL 10 – PENGAKHIRAN ======================
    add_pasal(doc, 10, "Pengakhiran Perjanjian")
    add_ayat_list(doc, [
        "Perjanjian ini dapat berakhir karena: (a) habisnya jangka waktu, (b) kesepakatan tertulis "
        "PARA PIHAK, (c) wanprestasi salah satu PIHAK yang tidak dipulihkan dalam 30 hari setelah teguran tertulis, "
        "(d) salah satu PIHAK mengalami pailit atau likuidasi.",

        "Pengakhiran Perjanjian tidak menghapus kewajiban PIHAK KEDUA untuk melunasi biaya Layanan "
        "yang telah timbul sampai tanggal pengakhiran.",

        "PARA PIHAK sepakat mengesampingkan ketentuan Pasal 1266 dan Pasal 1267 KUHPerdata sepanjang "
        "mengenai keharusan mendapatkan penetapan pengadilan untuk mengakhiri Perjanjian.",
    ])

    # ====================== PASAL 11 – FORCE MAJEURE ======================
    add_pasal(doc, 11, "Keadaan Kahar (Force Majeure)")
    add_ayat_list(doc, [
        "Keadaan Kahar adalah peristiwa di luar kendali PARA PIHAK yang tidak dapat diduga dan diatasi, "
        "termasuk namun tidak terbatas pada: bencana alam (gempa, banjir, kebakaran), perang, kerusuhan, "
        "pandemi, pemadaman listrik regional, gangguan internet nasional, peraturan pemerintah yang "
        "melarang pelaksanaan Perjanjian.",

        "PIHAK yang mengalami Force Majeure wajib memberitahukan PIHAK lainnya secara tertulis "
        "paling lambat 7 (tujuh) hari kalender setelah peristiwa terjadi.",

        "Selama masa Force Majeure, kewajiban PARA PIHAK ditangguhkan tanpa dikenakan sanksi.",

        "Jika Force Majeure berlangsung lebih dari 60 (enam puluh) hari, PARA PIHAK berhak mengakhiri "
        "Perjanjian dengan pemberitahuan tertulis.",
    ])

    # ====================== PASAL 12 – GANTI RUGI ======================
    add_pasal(doc, 12, "Batas Tanggung Jawab dan Ganti Rugi")
    add_ayat_list(doc, [
        "Tanggung jawab PIHAK PERTAMA atas seluruh klaim yang timbul dari Perjanjian ini dibatasi "
        "maksimum sebesar total biaya Layanan yang telah dibayar PIHAK KEDUA dalam 12 (dua belas) "
        "bulan terakhir sebelum peristiwa yang menimbulkan klaim.",

        "PIHAK PERTAMA tidak bertanggung jawab atas kerugian tidak langsung (indirect damages) seperti "
        "kehilangan keuntungan, kehilangan peluang bisnis, atau kerugian reputasi.",

        "PIHAK PERTAMA tidak bertanggung jawab atas kehilangan atau kerusakan DATA yang disebabkan "
        "oleh: kelalaian PENGGUNA, serangan siber pada infrastruktur PIHAK KEDUA, atau penggunaan "
        "APLIKASI yang tidak sesuai petunjuk.",
    ])

    # ====================== PASAL 13 – PENYELESAIAN SENGKETA ======================
    add_pasal(doc, 13, "Penyelesaian Sengketa")
    add_ayat_list(doc, [
        "Setiap perselisihan yang timbul dari Perjanjian ini akan diselesaikan terlebih dahulu secara "
        "musyawarah untuk mufakat dalam jangka waktu paling lambat 30 (tiga puluh) hari kalender.",

        "Apabila musyawarah tidak mencapai kesepakatan, PARA PIHAK sepakat menyelesaikan perselisihan "
        "melalui Pengadilan Negeri Bekasi.",

        "Perjanjian ini tunduk pada hukum Republik Indonesia.",
    ])

    # ====================== PASAL 14 – KETENTUAN LAIN ======================
    add_pasal(doc, 14, "Ketentuan Lain")
    add_ayat_list(doc, [
        "Setiap pemberitahuan tertulis dilakukan melalui email resmi kedua belah pihak dan dianggap "
        "sah setelah dikirim ke alamat email yang disepakati.",

        "Perubahan atau penambahan atas Perjanjian ini hanya sah bila dibuat secara tertulis dan "
        "ditandatangani oleh PARA PIHAK dalam bentuk Addendum yang menjadi bagian tak terpisahkan.",

        "Bila salah satu ketentuan Perjanjian ini dinyatakan tidak sah oleh pengadilan, ketentuan "
        "lainnya tetap berlaku sepanjang tidak menghilangkan maksud dan tujuan Perjanjian.",

        "Lampiran-lampiran Perjanjian ini merupakan bagian yang tidak terpisahkan.",
    ])

    # ====================== PASAL 15 – PENUTUP ======================
    add_pasal(doc, 15, "Penutup")
    add_para(doc,
             "Demikian Perjanjian Kerja Sama ini dibuat dalam rangkap 2 (dua) asli, masing-masing "
             "bermeterai cukup, mempunyai kekuatan hukum yang sama, dan ditandatangani oleh PARA PIHAK "
             "pada hari, tanggal, bulan, dan tahun sebagaimana disebutkan di awal Perjanjian ini.",
             align="justify")

    # ====================== TANDA TANGAN ======================
    doc.add_paragraph()
    doc.add_paragraph()
    table = doc.add_table(rows=6, cols=2)
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8)

    hdrs = table.rows[0].cells
    hdrs[0].text = "PIHAK PERTAMA"
    hdrs[1].text = "PIHAK KEDUA"
    for c in hdrs:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True

    r2 = table.rows[1].cells
    r2[0].text = "(Penyedia Aplikasi)"
    r2[1].text = "(Pengguna Aplikasi)"
    for c in r2:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # spacer rows for signature area
    for i in [2, 3, 4]:
        for c in table.rows[i].cells:
            c.text = ""

    # ttd row
    ttd = table.rows[5].cells
    ttd[0].text = "(_______________________________)\nNama:\nJabatan:"
    ttd[1].text = "(_______________________________)\nNama:\nJabatan:"
    for c in ttd:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ====================== LAMPIRAN ======================
    doc.add_page_break()
    add_heading(doc, "LAMPIRAN A", level=1, center=True)
    add_heading(doc, "SPESIFIKASI TEKNIS & BATAS LAYANAN", level=2, center=True)
    doc.add_paragraph()

    lampiran_table = doc.add_table(rows=8, cols=2)
    lampiran_table.style = "Light Grid"
    lamp_data = [
        ("Item", "Ketentuan"),
        ("Jumlah Pengguna (User)", "Sesuai kesepakatan (isi jumlah user)"),
        ("Akses", "Web-based, HTTPS"),
        ("Backup Data", "Harian, retensi 30 hari"),
        ("Uptime SLA", "99% per bulan"),
        ("Response Time Support", "Maksimal 1x24 jam pada Hari Kerja"),
        ("Update / Patch", "Terjadwal (notifikasi 24 jam sebelumnya)"),
        ("Format Export Data", "Excel (.xlsx), CSV"),
    ]
    for i, (k, v) in enumerate(lamp_data):
        row = lampiran_table.rows[i].cells
        row[0].text = k
        row[1].text = v
        if i == 0:
            for c in row:
                for p in c.paragraphs:
                    for run in p.runs: run.bold = True

    doc.add_paragraph()
    add_para(doc, "LAMPIRAN B", bold=True, align="center", size=13)
    add_para(doc, "SKEMA BIAYA LAYANAN", bold=True, align="center", size=12)
    add_para(doc, "")
    fee_table = doc.add_table(rows=5, cols=2)
    fee_table.style = "Light Grid"
    fee_data = [
        ("Item Biaya", "Nominal (Rp)"),
        ("Setup Fee (satu kali)", "___________________"),
        ("Subscription Fee — Bulanan", "___________________"),
        ("Custom Development (per manday)", "___________________"),
        ("Training On-site (per hari)", "___________________"),
    ]
    for i, (k, v) in enumerate(fee_data):
        row = fee_table.rows[i].cells
        row[0].text = k
        row[1].text = v
        if i == 0:
            for c in row:
                for p in c.paragraphs:
                    for run in p.runs: run.bold = True

    # ====================== SAVE ======================
    output = "/app/docs/PKS_Aplikasi_AR_Cargo_TMS.docx"
    doc.save(output)
    print(f"✅ Berhasil dibuat: {output}")


if __name__ == "__main__":
    build()
