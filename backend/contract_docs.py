"""
Generator dokumen Kontrak Kerja (PKWT) — .docx (python-docx) & .pdf (xhtml2pdf).
Pure-python, tanpa dependency system (aman untuk deploy Docker). Kop surat = logo AR CARGO
+ info perusahaan dari settings. Layout meniru template lampiran.
"""
import io
import os
import base64
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HARI = ["Minggu", "Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu"]
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
         "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

LOGO_PATH = os.path.join(os.path.dirname(__file__), "assets", "ar_cargo_logo.png")
BLANK = "_____________"


def _fmt_idr(n):
    try:
        return "Rp " + f"{int(round(float(n or 0))):,}".replace(",", ".")
    except Exception:
        return "Rp 0"


def _long_date(s):
    if not s:
        return "…………………"
    try:
        d = datetime.strptime(str(s)[:10], "%Y-%m-%d")
        return f"{d.day} {BULAN[d.month - 1]} {d.year}"
    except Exception:
        return str(s)


def _sign_parts(s):
    if not s:
        return {"hari": "…………", "tgl": "……", "bulan": "…………", "tahun": "………"}
    try:
        d = datetime.strptime(str(s)[:10], "%Y-%m-%d")
        return {"hari": HARI[(d.weekday() + 1) % 7],
                "tgl": d.day, "bulan": BULAN[d.month - 1], "tahun": d.year}
    except Exception:
        return {"hari": "…………", "tgl": "……", "bulan": "…………", "tahun": "………"}


def build_context(contract: dict, employee: dict, settings: dict) -> dict:
    c = contract or {}
    e = employee or {}
    s = settings or {}

    def pick(*vals):
        for v in vals:
            if v not in (None, "", 0):
                return v
        return None

    emp_name = pick(c.get("emp_name"), e.get("full_name"))
    emp_ktp = pick(c.get("emp_ktp"), e.get("nik"))
    emp_bp = pick(c.get("emp_birth_place"), e.get("tempat_lahir"))
    emp_bd = pick(c.get("emp_birth_date"), e.get("tanggal_lahir"))
    emp_addr = pick(c.get("emp_address"), e.get("alamat_ktp"), e.get("alamat_domisili"))

    ttl = ""
    if emp_bp or emp_bd:
        ttl = f"{emp_bp or ''}{', ' + _long_date(emp_bd) if emp_bd else ''}"

    return {
        "contract_type": c.get("contract_type") or "PKWT",
        "contract_no": c.get("contract_no"),
        "sign_place": c.get("sign_place") or "Kota Bekasi",
        "sign": _sign_parts(c.get("sign_date")),
        "company_name": c.get("company_name") or "PT. Abirama Mitra Pratama",
        "rep_name": c.get("rep_name"),
        "rep_position": c.get("rep_position"),
        "rep_address": c.get("rep_address"),
        "emp_name": emp_name,
        "emp_ktp": emp_ktp,
        "emp_ttl": ttl,
        "emp_address": emp_addr,
        "position": c.get("position"),
        "placement": c.get("placement"),
        "duration_text": c.get("duration_text"),
        "start_date": _long_date(c.get("start_date")),
        "end_date": _long_date(c.get("end_date")),
        "work_days_text": c.get("work_days_text"),
        "gaji_pokok": _fmt_idr(c.get("gaji_pokok")) if c.get("gaji_pokok") else None,
        "gaji_pokok_words": c.get("gaji_pokok_words"),
        "uang_makan": _fmt_idr(c.get("uang_makan")) if c.get("uang_makan") else None,
        "uang_makan_words": c.get("uang_makan_words"),
        "signer2_title": c.get("signer2_title") or "Karyawan",
        # Kop
        "kop_address": s.get("kop_address", ""),
        "kop_phone": s.get("kop_phone", ""),
        "kop_email": s.get("kop_email", ""),
        "kop_website": s.get("kop_website", ""),
        "kop_tagline": s.get("kop_tagline", ""),
    }


def _v(x):
    return x if x not in (None, "") else BLANK


# ---------------------------------------------------------------------------
# HTML (table-based, kompatibel xhtml2pdf → PDF)
# ---------------------------------------------------------------------------
def build_contract_html(ctx: dict) -> str:
    logo_uri = ""
    try:
        with open(LOGO_PATH, "rb") as f:
            logo_uri = "data:image/png;base64," + base64.b64encode(f.read()).decode()
    except Exception:
        logo_uri = ""

    sp = ctx["sign"]
    gaji = ctx["gaji_pokok"] or BLANK
    gaji_w = f' (<i>{ctx["gaji_pokok_words"]}</i>)' if ctx["gaji_pokok_words"] else ""
    makan = ctx["uang_makan"] or BLANK
    makan_w = f' (<i>{ctx["uang_makan_words"]}</i>)' if ctx["uang_makan_words"] else ""

    contact = []
    if ctx["kop_phone"]:
        contact.append(f'Telp: {ctx["kop_phone"]}')
    if ctx["kop_email"]:
        contact.append(f'Email: {ctx["kop_email"]}')
    if ctx["kop_website"]:
        contact.append(f'Website: {ctx["kop_website"]}')
    contact_line = " &middot; ".join(contact)

    logo_cell = f'<img src="{logo_uri}" width="120" />' if logo_uri else ""
    addr_html = f'<div class="kline">{ctx["kop_address"]}</div>' if ctx["kop_address"] else ""
    contact_html = f'<div class="kline">{contact_line}</div>' if contact_line else ""
    tagline_html = f'<div class="tagline">{ctx["kop_tagline"]}</div>' if ctx["kop_tagline"] else ""

    rep_sig = ctx["rep_name"] or "……………………"
    rep_pos = ctx["rep_position"] or "Direktur Utama"
    emp_sig = ctx["emp_name"] or "……………………"

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>
@page {{ size: a4 portrait; margin: 1.5cm 2cm; }}
body {{ font-family: "Times-Roman", serif; font-size: 11pt; color: #111; line-height: 1.4; }}
.cname {{ font-size: 15pt; font-weight: bold; }}
.kline {{ font-size: 8.5pt; color: #333; }}
.tagline {{ font-size: 8pt; font-style: italic; color: #555; }}
hr {{ border: none; border-bottom: 2px solid #1e3a8a; margin: 3px 0 10px 0; }}
.title {{ text-align: center; font-weight: bold; font-size: 13pt; margin-top: 4px; }}
.subno {{ text-align: center; margin-bottom: 4px; }}
p {{ text-align: justify; margin: 5px 0; }}
.ph {{ text-align: center; font-weight: bold; margin: 1px 0; -pdf-keep-with-next: true; }}
.pasal {{ margin-top: 8px; }}
ol {{ margin: 2px 0 2px 16px; padding: 0; }}
li {{ text-align: justify; margin-bottom: 2px; }}
td {{ vertical-align: top; }}
.nm {{ font-weight: bold; }}
</style></head><body>

<table width="100%"><tr>
  <td width="24%">{logo_cell}</td>
  <td align="center">
    <div class="cname">{ctx["company_name"]}</div>
    {addr_html}{contact_html}{tagline_html}
  </td>
</tr></table>
<hr/>

<div class="title">PERJANJIAN KERJA WAKTU TERTENTU ({ctx["contract_type"]})</div>
<div class="subno">No: {_v(ctx["contract_no"])}</div>

<p>Pada hari ini, <b>{sp["hari"]}</b>, tanggal <b>{sp["tgl"]}</b> bulan <b>{sp["bulan"]}</b> tahun <b>{sp["tahun"]}</b>, bertempat di {ctx["sign_place"]}, yang bertanda tangan di bawah ini :</p>

<table width="100%">
  <tr><td width="26%">Nama</td><td width="3%">:</td><td>{_v(ctx["rep_name"])}</td></tr>
  <tr><td>Jabatan</td><td>:</td><td>{_v(ctx["rep_position"])}</td></tr>
  <tr><td>Alamat</td><td>:</td><td>{_v(ctx["rep_address"])}</td></tr>
</table>
<p>Dalam hal ini bertindak untuk dan atas nama <b>{ctx["company_name"]}</b>, selanjutnya disebut sebagai <b>PIHAK PERTAMA</b>.</p>

<table width="100%">
  <tr><td width="26%">Nama</td><td width="3%">:</td><td>{_v(ctx["emp_name"])}</td></tr>
  <tr><td>No. KTP</td><td>:</td><td>{_v(ctx["emp_ktp"])}</td></tr>
  <tr><td>Tempat, Tgl Lahir</td><td>:</td><td>{_v(ctx["emp_ttl"])}</td></tr>
  <tr><td>Alamat</td><td>:</td><td>{_v(ctx["emp_address"])}</td></tr>
</table>
<p>Dalam hal ini bertindak untuk dan atas nama diri sendiri, selanjutnya disebut sebagai <b>PIHAK KEDUA</b>.</p>

<p>PIHAK PERTAMA dan PIHAK KEDUA secara bersama-sama disebut "Para Pihak".</p>
<p>Para Pihak sepakat untuk mengikatkan diri dalam Perjanjian Kerja Waktu Tertentu (selanjutnya disebut "Perjanjian") dengan syarat-syarat dan ketentuan sebagai berikut :</p>

<div class="pasal"><div class="ph">Pasal 1</div><div class="ph">TUGAS DAN PENEMPATAN</div>
<ol>
<li>PIHAK PERTAMA menerima dan mempekerjakan PIHAK KEDUA sebagai <b>{_v(ctx["position"])}</b>.</li>
<li>PIHAK KEDUA akan ditempatkan di {_v(ctx["placement"])}.</li>
<li>PIHAK PERTAMA berhak memindahkan, merotasi, atau menugaskan PIHAK KEDUA ke divisi atau lokasi kerja lain sesuai dengan kebutuhan operasional perusahaan.</li>
</ol></div>

<div class="pasal"><div class="ph">Pasal 2</div><div class="ph">JANGKA WAKTU PERJANJIAN</div>
<p>Perjanjian ini berlaku selama <b>{_v(ctx["duration_text"])}</b>, terhitung efektif sejak tanggal <b>{ctx["start_date"]}</b> dan akan berakhir dengan sendirinya pada tanggal <b>{ctx["end_date"]}</b>.</p></div>

<div class="pasal"><div class="ph">Pasal 3</div><div class="ph">WAKTU KERJA DAN TATA TERTIB</div>
<ol>
<li>Waktu kerja dan jam kerja operasional yang berlaku bagi PIHAK KEDUA disesuaikan dengan ritme kebutuhan operasional perusahaan, yaitu <b>{_v(ctx["work_days_text"])}</b>.</li>
<li>Pengaturan lebih lanjut mengenai rincian jam masuk, jam pulang, dan waktu istirahat akan diatur secara terpisah dalam Peraturan Perusahaan dan/atau Standar Operasional Prosedur (SOP) yang ditetapkan oleh PIHAK PERTAMA.</li>
<li>PIHAK KEDUA wajib tunduk dan patuh pada Peraturan Perusahaan, Standar Operasional Prosedur (SOP), serta kebijakan tata tertib kehadiran yang berlaku di lingkungan {ctx["company_name"]}.</li>
</ol></div>

<div class="pasal"><div class="ph">Pasal 4</div><div class="ph">UPAH, TUNJANGAN, DAN KERJA LEMBUR</div>
<ol>
<li>Selama masa Perjanjian ini, PIHAK PERTAMA akan memberikan kompensasi kepada PIHAK KEDUA dengan rincian sebagai berikut :
<ol style="list-style-type: lower-alpha">
<li>Gaji Pokok : Sebesar <b>{gaji}</b>{gaji_w} nett/bulan.</li>
<li>Tunjangan Tidak Tetap (Uang Makan): Sebesar <b>{makan}</b>{makan_w} / hari kehadiran kerja.</li>
</ol></li>
<li>Uang makan sebagaimana diatur pada ayat (1) huruf b bersifat fluktuatif dan hanya dibayarkan apabila PIHAK KEDUA hadir secara fisik serta menyelesaikan waktu kerja reguler pada hari tersebut. Apabila PIHAK KEDUA mangkir, sakit, cuti, atau izin, maka uang makan tidak dibayarkan.</li>
<li>Upah sebagaimana dimaksud pada ayat (1) dibayarkan pada tanggal terakhir setiap bulannya melalui transfer ke rekening bank milik PIHAK KEDUA.</li>
<li>Apabila terdapat kebutuhan operasional yang mendesak, PIHAK PERTAMA dapat menugaskan PIHAK KEDUA untuk melakukan Kerja Lembur. Pelaksanaan Kerja Lembur wajib didasarkan pada Surat Perintah Lembur (SPL) resmi yang disetujui dan ditandatangani terlebih dahulu oleh atasan langsung PIHAK KEDUA. Tanpa adanya SPL tertulis, segala bentuk perpanjangan waktu kerja (kehadiran di luar jam kerja operasional atas inisiatif pribadi) tidak diakui dan tidak sah dibayarkan sebagai lembur.</li>
<li>Perhitungan dan pembayaran upah Kerja Lembur yang disetujui melalui SPL akan dilaksanakan sesuai dengan kebijakan internal yang sejalan dengan peraturan perundang-undangan di bidang ketenagakerjaan yang berlaku.</li>
</ol></div>

<div class="pasal"><div class="ph">Pasal 5</div><div class="ph">FASILITAS DAN KERAHASIAAN PERUSAHAAN</div>
<ol>
<li>PIHAK PERTAMA dapat meminjamkan fasilitas atau perlengkapan kerja kepada PIHAK KEDUA. Seluruh aset tersebut tetap menjadi milik sah PIHAK PERTAMA dan wajib dikembalikan dalam keadaan baik saat Perjanjian ini berakhir.</li>
<li>PIHAK KEDUA diwajibkan untuk menjaga kerahasiaan seluruh data operasional, administrasi, sistem kerja, serta informasi konfidensial lainnya milik PIHAK PERTAMA, baik selama Perjanjian ini berlangsung maupun setelah berakhirnya hubungan kerja.</li>
</ol></div>

<div class="pasal"><div class="ph">Pasal 6</div><div class="ph">PENGAKHIRAN PERJANJIAN</div>
<ol>
<li>Pemutusan hubungan kerja dapat terjadi sebelum berakhirnya jangka waktu Perjanjian apabila PIHAK KEDUA melakukan pelanggaran berat, melanggar tata tertib perusahaan, atau tidak mampu memenuhi standar kinerja yang ditetapkan oleh PIHAK PERTAMA.</li>
<li>Pemenuhan hak-hak normatif PIHAK KEDUA akibat berakhirnya Perjanjian ini akan diselesaikan dengan mengacu pada ketentuan perundang-undangan di bidang ketenagakerjaan yang berlaku di Republik Indonesia.</li>
</ol></div>

<div class="pasal"><div class="ph">Pasal 7</div><div class="ph">PENUTUP</div>
<ol>
<li>Perjanjian ini dibuat secara sadar, dibaca, dimengerti, dan disepakati bersama. Khusus mengenai besaran Upah sebagaimana diatur dalam Pasal 4, Para Pihak menyadari bahwa PIHAK PERTAMA diklasifikasikan sebagai Usaha Kecil berdasarkan regulasi yang berlaku, sehingga besaran upah ditetapkan murni berdasarkan kesepakatan Para Pihak.</li>
<li>Dibuat rangkap 2 (dua) yang masing-masing dibubuhi meterai secukupnya dan memiliki kekuatan hukum yang sama.</li>
</ol></div>

<table width="100%" style="margin-top:34px">
<tr>
  <td width="50%" align="center">PIHAK PERTAMA<br/>{ctx["company_name"]}<br/><br/><br/><br/><span class="nm">( {rep_sig} )</span><br/>{rep_pos}</td>
  <td width="50%" align="center">PIHAK KEDUA<br/><br/><br/><br/><br/><span class="nm">( {emp_sig} )</span><br/>{ctx["signer2_title"]}</td>
</tr>
</table>

</body></html>"""


def build_contract_pdf(ctx: dict) -> bytes:
    from xhtml2pdf import pisa
    html = build_contract_html(ctx)
    out = io.BytesIO()
    pisa.CreatePDF(html, dest=out, encoding="utf-8")
    return out.getvalue()


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------
def _set_cell_border_bottom(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "double")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), "1e3a8a")
    borders.append(bottom)
    tc_pr.append(borders)


def _p(doc, text="", bold=False, align=None, size=None, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return p


def _num_item(doc, runs, num, indent=0.8, alpha=False):
    """Numbered list item with manual numbering. runs = list of (text, bold, italic)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    prefix = f"{num}. "
    p.add_run(prefix)
    for (t, b, i) in runs:
        r = p.add_run(t)
        r.bold = b
        r.italic = i
    return p


def build_contract_docx(ctx: dict) -> bytes:
    doc = Document()
    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11.5)
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # --- Kop surat: 2-col borderless table (logo | info) ---
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    left, right = tbl.rows[0].cells
    left.width = Cm(4.5)
    right.width = Cm(12.5)
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(LOGO_PATH):
        try:
            lp.add_run().add_picture(LOGO_PATH, width=Cm(4.0))
        except Exception:
            pass
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = rp.add_run(ctx["company_name"])
    cr.bold = True
    cr.font.size = Pt(15)
    if ctx["kop_address"]:
        pa = right.add_paragraph()
        pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ra = pa.add_run(ctx["kop_address"]); ra.font.size = Pt(9)
    contact = []
    if ctx["kop_phone"]:
        contact.append(f'Telp: {ctx["kop_phone"]}')
    if ctx["kop_email"]:
        contact.append(f'Email: {ctx["kop_email"]}')
    if ctx["kop_website"]:
        contact.append(f'Website: {ctx["kop_website"]}')
    if contact:
        pc = right.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = pc.add_run(" · ".join(contact)); rc.font.size = Pt(9)
    if ctx["kop_tagline"]:
        pt = right.add_paragraph(); pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rt = pt.add_run(ctx["kop_tagline"]); rt.italic = True; rt.font.size = Pt(8.5)
    _set_cell_border_bottom(left)
    _set_cell_border_bottom(right)

    doc.add_paragraph()
    # Title
    _p(doc, f'PERJANJIAN KERJA WAKTU TERTENTU ({ctx["contract_type"]})', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13.5)
    _p(doc, f'No: {_v(ctx["contract_no"])}', align=WD_ALIGN_PARAGRAPH.CENTER)

    sp = ctx["sign"]
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Pada hari ini, ")
    p.add_run(str(sp["hari"])).bold = True
    p.add_run(", tanggal ")
    p.add_run(str(sp["tgl"])).bold = True
    p.add_run(" bulan ")
    p.add_run(str(sp["bulan"])).bold = True
    p.add_run(" tahun ")
    p.add_run(str(sp["tahun"])).bold = True
    p.add_run(f', bertempat di {ctx["sign_place"]}, yang bertanda tangan di bawah ini :')

    def party(rows):
        t = doc.add_table(rows=len(rows), cols=3)
        t.allow_autofit = True
        for i, (k, v) in enumerate(rows):
            c0, c1, c2 = t.rows[i].cells
            c0.width = Cm(4.2); c1.width = Cm(0.5); c2.width = Cm(12)
            c0.paragraphs[0].add_run(k)
            c1.paragraphs[0].add_run(":")
            c2.paragraphs[0].add_run(v)

    party([("Nama", _v(ctx["rep_name"])), ("Jabatan", _v(ctx["rep_position"])), ("Alamat", _v(ctx["rep_address"]))])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Dalam hal ini bertindak untuk dan atas nama ")
    p.add_run(ctx["company_name"]).bold = True
    p.add_run(", selanjutnya disebut sebagai ")
    p.add_run("PIHAK PERTAMA").bold = True
    p.add_run(".")

    party([("Nama", _v(ctx["emp_name"])), ("No. KTP", _v(ctx["emp_ktp"])),
           ("Tempat, Tgl Lahir", _v(ctx["emp_ttl"])), ("Alamat", _v(ctx["emp_address"]))])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Dalam hal ini bertindak untuk dan atas nama diri sendiri, selanjutnya disebut sebagai ")
    p.add_run("PIHAK KEDUA").bold = True
    p.add_run(".")

    _p(doc, 'PIHAK PERTAMA dan PIHAK KEDUA secara bersama-sama disebut "Para Pihak".', align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _p(doc, 'Para Pihak sepakat untuk mengikatkan diri dalam Perjanjian Kerja Waktu Tertentu (selanjutnya disebut "Perjanjian") dengan syarat-syarat dan ketentuan sebagai berikut :', align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def pasal(title, sub):
        _p(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _p(doc, sub, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Pasal 1
    pasal("Pasal 1", "TUGAS DAN PENEMPATAN")
    _num_item(doc, [("PIHAK PERTAMA menerima dan mempekerjakan PIHAK KEDUA sebagai ", False, False), (_v(ctx["position"]), True, False), (".", False, False)], 1)
    _num_item(doc, [(f'PIHAK KEDUA akan ditempatkan di {_v(ctx["placement"])}.', False, False)], 2)
    _num_item(doc, [("PIHAK PERTAMA berhak memindahkan, merotasi, atau menugaskan PIHAK KEDUA ke divisi atau lokasi kerja lain sesuai dengan kebutuhan operasional perusahaan.", False, False)], 3)

    # Pasal 2
    pasal("Pasal 2", "JANGKA WAKTU PERJANJIAN")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("Perjanjian ini berlaku selama ")
    p.add_run(_v(ctx["duration_text"])).bold = True
    p.add_run(", terhitung efektif sejak tanggal ")
    p.add_run(ctx["start_date"]).bold = True
    p.add_run(" dan akan berakhir dengan sendirinya pada tanggal ")
    p.add_run(ctx["end_date"]).bold = True
    p.add_run(".")

    # Pasal 3
    pasal("Pasal 3", "WAKTU KERJA DAN TATA TERTIB")
    _num_item(doc, [("Waktu kerja dan jam kerja operasional yang berlaku bagi PIHAK KEDUA disesuaikan dengan ritme kebutuhan operasional perusahaan, yaitu ", False, False), (_v(ctx["work_days_text"]), True, False), (".", False, False)], 1)
    _num_item(doc, [("Pengaturan lebih lanjut mengenai rincian jam masuk, jam pulang, dan waktu istirahat akan diatur secara terpisah dalam Peraturan Perusahaan dan/atau Standar Operasional Prosedur (SOP) yang ditetapkan oleh PIHAK PERTAMA.", False, False)], 2)
    _num_item(doc, [(f'PIHAK KEDUA wajib tunduk dan patuh pada Peraturan Perusahaan, Standar Operasional Prosedur (SOP), serta kebijakan tata tertib kehadiran yang berlaku di lingkungan {ctx["company_name"]}.', False, False)], 3)

    # Pasal 4
    pasal("Pasal 4", "UPAH, TUNJANGAN, DAN KERJA LEMBUR")
    _num_item(doc, [("Selama masa Perjanjian ini, PIHAK PERTAMA akan memberikan kompensasi kepada PIHAK KEDUA dengan rincian sebagai berikut :", False, False)], 1)
    ga = doc.add_paragraph(); ga.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; ga.paragraph_format.left_indent = Cm(1.6)
    ga.add_run("a. Gaji Pokok : Sebesar ")
    ga.add_run(ctx["gaji_pokok"] or BLANK).bold = True
    if ctx["gaji_pokok_words"]:
        ga.add_run(" ("); ga.add_run(ctx["gaji_pokok_words"]).italic = True; ga.add_run(")")
    ga.add_run(" nett/bulan.")
    gb = doc.add_paragraph(); gb.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; gb.paragraph_format.left_indent = Cm(1.6)
    gb.add_run("b. Tunjangan Tidak Tetap (Uang Makan): Sebesar ")
    gb.add_run(ctx["uang_makan"] or BLANK).bold = True
    if ctx["uang_makan_words"]:
        gb.add_run(" ("); gb.add_run(ctx["uang_makan_words"]).italic = True; gb.add_run(")")
    gb.add_run(" / hari kehadiran kerja.")
    _num_item(doc, [("Uang makan sebagaimana diatur pada ayat (1) huruf b bersifat fluktuatif dan hanya dibayarkan apabila PIHAK KEDUA hadir secara fisik serta menyelesaikan waktu kerja reguler pada hari tersebut. Apabila PIHAK KEDUA mangkir, sakit, cuti, atau izin, maka uang makan tidak dibayarkan.", False, False)], 2)
    _num_item(doc, [("Upah sebagaimana dimaksud pada ayat (1) dibayarkan pada tanggal terakhir setiap bulannya melalui transfer ke rekening bank milik PIHAK KEDUA.", False, False)], 3)
    _num_item(doc, [("Apabila terdapat kebutuhan operasional yang mendesak, PIHAK PERTAMA dapat menugaskan PIHAK KEDUA untuk melakukan Kerja Lembur. Pelaksanaan Kerja Lembur wajib didasarkan pada Surat Perintah Lembur (SPL) resmi yang disetujui dan ditandatangani terlebih dahulu oleh atasan langsung PIHAK KEDUA. Tanpa adanya SPL tertulis, segala bentuk perpanjangan waktu kerja (kehadiran di luar jam kerja operasional atas inisiatif pribadi) tidak diakui dan tidak sah dibayarkan sebagai lembur.", False, False)], 4)
    _num_item(doc, [("Perhitungan dan pembayaran upah Kerja Lembur yang disetujui melalui SPL akan dilaksanakan sesuai dengan kebijakan internal yang sejalan dengan peraturan perundang-undangan di bidang ketenagakerjaan yang berlaku.", False, False)], 5)

    # Pasal 5
    pasal("Pasal 5", "FASILITAS DAN KERAHASIAAN PERUSAHAAN")
    _num_item(doc, [("PIHAK PERTAMA dapat meminjamkan fasilitas atau perlengkapan kerja kepada PIHAK KEDUA. Seluruh aset tersebut tetap menjadi milik sah PIHAK PERTAMA dan wajib dikembalikan dalam keadaan baik saat Perjanjian ini berakhir.", False, False)], 1)
    _num_item(doc, [("PIHAK KEDUA diwajibkan untuk menjaga kerahasiaan seluruh data operasional, administrasi, sistem kerja, serta informasi konfidensial lainnya milik PIHAK PERTAMA, baik selama Perjanjian ini berlangsung maupun setelah berakhirnya hubungan kerja.", False, False)], 2)

    # Pasal 6
    pasal("Pasal 6", "PENGAKHIRAN PERJANJIAN")
    _num_item(doc, [("Pemutusan hubungan kerja dapat terjadi sebelum berakhirnya jangka waktu Perjanjian apabila PIHAK KEDUA melakukan pelanggaran berat, melanggar tata tertib perusahaan, atau tidak mampu memenuhi standar kinerja yang ditetapkan oleh PIHAK PERTAMA.", False, False)], 1)
    _num_item(doc, [("Pemenuhan hak-hak normatif PIHAK KEDUA akibat berakhirnya Perjanjian ini akan diselesaikan dengan mengacu pada ketentuan perundang-undangan di bidang ketenagakerjaan yang berlaku di Republik Indonesia.", False, False)], 2)

    # Pasal 7
    pasal("Pasal 7", "PENUTUP")
    _num_item(doc, [("Perjanjian ini dibuat secara sadar, dibaca, dimengerti, dan disepakati bersama. Khusus mengenai besaran Upah sebagaimana diatur dalam Pasal 4, Para Pihak menyadari bahwa PIHAK PERTAMA diklasifikasikan sebagai Usaha Kecil berdasarkan regulasi yang berlaku, sehingga besaran upah ditetapkan murni berdasarkan kesepakatan Para Pihak.", False, False)], 1)
    _num_item(doc, [("Dibuat rangkap 2 (dua) yang masing-masing dibubuhi meterai secukupnya dan memiliki kekuatan hukum yang sama.", False, False)], 2)

    # Signatures
    doc.add_paragraph()
    sig = doc.add_table(rows=1, cols=2)
    c0, c1 = sig.rows[0].cells
    for cell, lines in [
        (c0, ["PIHAK PERTAMA", ctx["company_name"], "", "", "", f'( {ctx["rep_name"] or "……………………"} )', ctx["rep_position"] or "Direktur Utama"]),
        (c1, ["PIHAK KEDUA", "", "", "", "", f'( {ctx["emp_name"] or "……………………"} )', ctx["signer2_title"]]),
    ]:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run(lines[0])
        for ln in lines[1:]:
            pp = cell.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = pp.add_run(ln)
            if ln.startswith("("):
                run.bold = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
