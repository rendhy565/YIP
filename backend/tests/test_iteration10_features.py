"""
Iteration-10 backend regression tests.

Covers:
- Tax options no longer include PPN 12 (only 0/1/11)
- OrderIn polis_amount + insurance_amount (no more insurance_policy)
- PATCH /api/orders/{oid}/extended polis_amount persists
- POST /api/orders/{oid}/receive auto-creates awb_costs for SPSI/Sewa Forklift,
  persists received_polis_amount/received_asuransi_amount, no-duplicate semantics
- POST /api/orders/bulk-revert-to-open (multi + empty)
- Per-Issuer bank settings present in /api/settings
- PUT /api/settings/{key} accepts body without key_name (Optional key_name regression)
- POST /api/settings/invoice-logo admin-only, persists invoice_logo_b64
- GET /api/exports/invoice/{id} formatting checks (INVOICE banner, header row, summary labels, bank box)
- GET /api/exports/report-ar?invoice_ids=... filter
- POST /api/cash-advances + /settle awb_breakdown happy path, mismatch -> 400,
  partial breakdown + refund happy path
"""
import os
import io
import time
import base64
from typing import Optional, List, Dict, Tuple

import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
assert BASE_URL, "REACT_APP_BACKEND_URL must be set"

API = f"{BASE_URL}/api"

TINY_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


# ------------------------- Fixtures -----------------------------------------
def _login(username: str, password: str) -> str:
    r = requests.post(f"{API}/auth/login", json={"username": username, "password": password}, timeout=20)
    assert r.status_code == 200, f"login {username} failed: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="session")
def admin_token():
    return _login("admin", "admin123")


@pytest.fixture(scope="session")
def operation_token():
    return _login("operation", "operation123")


@pytest.fixture(scope="session")
def finance_token():
    return _login("finance", "finance123")


def _headers(tok: str) -> Dict[str, str]:
    return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}


@pytest.fixture(scope="session")
def seed_customer(operation_token) -> str:
    # Try to reuse a TEST_ customer or create one
    r = requests.get(f"{API}/customers", headers=_headers(operation_token), timeout=20)
    assert r.status_code == 200, r.text
    items = r.json()
    for it in items:
        if it.get("name", "").startswith("TEST_IT10_CUST"):
            return it["id"]
    payload = {
        "name": f"TEST_IT10_CUST_{int(time.time())}",
        "phone": "0812",
        "address": "Jakarta",
        "npwp": "",
        "is_corporate": True,
    }
    r = requests.post(f"{API}/customers", headers=_headers(operation_token), json=payload, timeout=20)
    assert r.status_code in (200, 201), r.text
    return r.json()["id"]


def _create_order(operation_token: str, customer_id: str, *,
                  insurance: bool = False, polis: float = 0, asuransi: float = 0) -> dict:
    payload = {
        "customer_id": customer_id,
        "item_description": "TEST_IT10",
        "goods_content": "Box",
        "origin": "Jakarta",
        "destination": "Surabaya",
        "tariff_per_kg": 1000,
        "qty_kg": 10,
        "tax_percentage": 11,
        "company_name": "PT. Abirama Mitra Pradana",
        "service": "AIR",
        "via": "U",
        "insurance_enabled": insurance,
        "polis_amount": polis if insurance else None,
        "insurance_amount": asuransi if insurance else None,
    }
    r = requests.post(f"{API}/orders", headers=_headers(operation_token), json=payload, timeout=20)
    assert r.status_code in (200, 201), f"create_order failed: {r.status_code} {r.text}"
    return r.json()


# ------------------------- Tax options --------------------------------------
class TestTaxOptions:
    def test_tax_options_no_ppn12(self, operation_token):
        r = requests.get(f"{API}/settings/tax-options", headers=_headers(operation_token), timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        labels = [d["label"] for d in data]
        values = sorted({float(d["value"]) for d in data})
        assert 12.0 not in values, f"PPN 12 should be removed; got values: {values}, labels: {labels}"
        # Must contain exactly 0,1,11
        assert set(values) == {0.0, 1.0, 11.0}, f"expected exactly 0/1/11, got {values} (labels={labels})"


# ------------------------- OrderIn polis_amount -----------------------------
class TestOrderPolisAmount:
    def test_create_order_with_polis_and_insurance(self, operation_token, seed_customer):
        data = _create_order(operation_token, seed_customer,
                             insurance=True, polis=20000, asuransi=150600)
        assert data.get("polis_amount") == 20000, f"polis_amount missing or wrong: {data.get('polis_amount')}"
        assert data.get("insurance_amount") == 150600, f"insurance_amount missing or wrong: {data.get('insurance_amount')}"
        # legacy field must not be present
        assert "insurance_policy" not in data, "insurance_policy should be removed"
        assert data.get("insurance_enabled") is True

    def test_patch_extended_polis_amount(self, operation_token, seed_customer):
        # Need an Open order
        order = _create_order(operation_token, seed_customer, insurance=True, polis=10000, asuransi=50000)
        oid = order["id"]
        r = requests.patch(f"{API}/orders/{oid}/extended",
                           headers=_headers(operation_token),
                           json={"polis_amount": 50000}, timeout=20)
        assert r.status_code == 200, r.text
        assert r.json().get("polis_amount") == 50000


# ------------------------- Receive + auto-cost ------------------------------
class TestReceiveAutoCost:
    @pytest.fixture(scope="class")
    def receive_order(self, operation_token, seed_customer) -> str:
        o = _create_order(operation_token, seed_customer, insurance=True, polis=10000, asuransi=50000)
        return o["id"]

    def _receive(self, operation_token, oid, **overrides):
        body = {
            "received_by_name": "TEST_RECEIVER",
            "polis_amount": 25000,
            "asuransi_amount": 100000,
            "spsi_enabled": True,
            "spsi_amount": 75000,
            "sewa_forklift_enabled": True,
            "sewa_forklift_amount": 50000,
            "work_type": "Kirim Barang",
            "mode": "Trucking",
        }
        body.update(overrides)
        return requests.post(f"{API}/orders/{oid}/receive",
                             headers=_headers(operation_token),
                             json=body, timeout=20)

    def test_receive_creates_spsi_and_forklift_costs(self, operation_token, receive_order):
        oid = receive_order
        r = self._receive(operation_token, oid)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("received_polis_amount") == 25000
        assert body.get("received_asuransi_amount") == 100000

        # Verify awb_costs has SPSI + Sewa Forklift
        rc = requests.get(f"{API}/costs?awb_id={oid}", headers=_headers(operation_token), timeout=20)
        assert rc.status_code == 200, rc.text
        costs = rc.json()
        receive_costs = [c for c in costs if c.get("source") == "receive"]
        labels = sorted(c.get("sub_label") for c in receive_costs)
        assert "SPSI" in labels, f"SPSI missing in costs: {receive_costs}"
        assert "Sewa Forklift" in labels, f"Sewa Forklift missing in costs: {receive_costs}"
        # cost_category Operational on both
        for c in receive_costs:
            assert c.get("cost_category") == "Operational", c
        spsi = next(c for c in receive_costs if c["sub_label"] == "SPSI")
        forklift = next(c for c in receive_costs if c["sub_label"] == "Sewa Forklift")
        assert float(spsi["amount"]) == 75000
        assert float(forklift["amount"]) == 50000

    def test_receive_replaces_previous_spsi_cost_no_duplicate(self, operation_token, receive_order):
        oid = receive_order
        # Second receive with different SPSI amount
        r = self._receive(operation_token, oid, spsi_amount=80000)
        assert r.status_code == 200, r.text
        rc = requests.get(f"{API}/costs?awb_id={oid}", headers=_headers(operation_token), timeout=20)
        assert rc.status_code == 200
        receive_costs = [c for c in rc.json() if c.get("source") == "receive"]
        spsi_costs = [c for c in receive_costs if c.get("sub_label") == "SPSI"]
        forklift_costs = [c for c in receive_costs if c.get("sub_label") == "Sewa Forklift"]
        assert len(spsi_costs) == 1, f"expected 1 SPSI, got {len(spsi_costs)}: {spsi_costs}"
        assert len(forklift_costs) == 1, f"expected 1 Forklift, got {len(forklift_costs)}: {forklift_costs}"
        assert float(spsi_costs[0]["amount"]) == 80000

    def test_receive_spsi_disabled_no_cost(self, operation_token, seed_customer):
        o = _create_order(operation_token, seed_customer)
        oid = o["id"]
        r = requests.post(f"{API}/orders/{oid}/receive",
                          headers=_headers(operation_token),
                          json={
                              "received_by_name": "X",
                              "spsi_enabled": False,
                              "sewa_forklift_enabled": False,
                              "work_type": "Kirim Barang",
                              "mode": "Air",
                          }, timeout=20)
        assert r.status_code == 200, r.text
        rc = requests.get(f"{API}/costs?awb_id={oid}", headers=_headers(operation_token), timeout=20)
        assert rc.status_code == 200
        receive_costs = [c for c in rc.json() if c.get("source") == "receive"]
        assert all(c.get("sub_label") != "SPSI" for c in receive_costs), receive_costs


# ------------------------- Bulk revert --------------------------------------
class TestBulkRevert:
    def test_bulk_revert_to_open_multi(self, operation_token, finance_token, seed_customer):
        # Create N orders, receive + reconcile them to Reconciled, then revert.
        oids = []
        for _ in range(2):
            o = _create_order(operation_token, seed_customer)
            oid = o["id"]
            r = requests.post(f"{API}/orders/{oid}/receive",
                              headers=_headers(operation_token),
                              json={"received_by_name": "T",
                                    "spsi_enabled": True, "spsi_amount": 10000,
                                    "work_type": "Kirim Barang", "mode": "Air"}, timeout=20)
            assert r.status_code == 200, r.text
            oids.append(oid)
        # Reconcile
        r = requests.post(f"{API}/orders/reconcile",
                          headers=_headers(operation_token),
                          json={"awb_ids": oids}, timeout=20)
        assert r.status_code == 200, r.text
        # Bulk revert
        r = requests.post(f"{API}/orders/bulk-revert-to-open",
                          headers=_headers(operation_token),
                          json={"awb_ids": oids}, timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("updated") == 2, body
        # Verify status Open
        for oid in oids:
            r2 = requests.get(f"{API}/orders/{oid}", headers=_headers(operation_token), timeout=20)
            assert r2.status_code == 200
            assert r2.json()["status"] == "Open"

    def test_bulk_revert_empty(self, operation_token):
        r = requests.post(f"{API}/orders/bulk-revert-to-open",
                          headers=_headers(operation_token),
                          json={"awb_ids": []}, timeout=20)
        assert r.status_code == 200, r.text
        assert r.json().get("updated") == 0


# ------------------------- Per-Issuer bank settings -------------------------
class TestPerIssuerBankSettings:
    def test_bank_settings_present(self, operation_token):
        r = requests.get(f"{API}/settings", headers=_headers(operation_token), timeout=20)
        assert r.status_code == 200
        keys = {s["key_name"] for s in r.json()}
        expected = {
            "bank_name_PT. Abirama Mitra Pradana",
            "bank_name_PT. Abirama Mitra Pratama",
            "bank_holder_PT. Abirama Mitra Pradana",
            "bank_holder_PT. Abirama Mitra Pratama",
            "bank_account_PT. Abirama Mitra Pradana",
            "bank_account_PT. Abirama Mitra Pratama",
        }
        missing = expected - keys
        assert not missing, f"Missing bank keys: {missing}"

    def test_put_settings_without_keyname_in_body(self, operation_token):
        # Regression: SettingIn.key_name should be Optional
        key = "bank_account_PT. Abirama Mitra Pradana"
        r = requests.put(f"{API}/settings/{key}",
                         headers=_headers(operation_token),
                         json={"key_value": "1234567890"}, timeout=20)
        assert r.status_code == 200, r.text
        # Verify
        r2 = requests.get(f"{API}/settings", headers=_headers(operation_token), timeout=20)
        m = {s["key_name"]: s["key_value"] for s in r2.json()}
        assert m.get(key) == "1234567890"


# ------------------------- Invoice logo upload ------------------------------
class TestInvoiceLogo:
    def test_upload_logo_admin(self, admin_token, operation_token):
        r = requests.post(f"{API}/settings/invoice-logo",
                          headers=_headers(admin_token),
                          json={"image_b64": TINY_PNG_B64}, timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("ok") is True
        assert isinstance(body.get("size_bytes"), int) and body["size_bytes"] > 0
        # verify settings includes invoice_logo_b64 non-empty
        r2 = requests.get(f"{API}/settings", headers=_headers(operation_token), timeout=20)
        m = {s["key_name"]: s["key_value"] for s in r2.json()}
        assert m.get("invoice_logo_b64"), "invoice_logo_b64 should be persisted non-empty"

    def test_upload_logo_non_admin_forbidden(self, operation_token):
        r = requests.post(f"{API}/settings/invoice-logo",
                          headers=_headers(operation_token),
                          json={"image_b64": TINY_PNG_B64}, timeout=20)
        assert r.status_code == 403, f"expected 403, got {r.status_code}: {r.text}"


# ------------------------- Invoice Excel & Report AR exports ----------------
def _ensure_invoice(operation_token: str, finance_token: str, customer_id: str) -> Tuple[str, str]:
    """Return (invoice_id, awb_id) after creating & reconciling a new order."""
    o = _create_order(operation_token, customer_id, insurance=True, polis=15000, asuransi=75000)
    oid = o["id"]
    r = requests.post(f"{API}/orders/{oid}/receive",
                      headers=_headers(operation_token),
                      json={"received_by_name": "T", "polis_amount": 15000, "asuransi_amount": 75000,
                            "spsi_enabled": True, "spsi_amount": 30000,
                            "sewa_forklift_enabled": True, "sewa_forklift_amount": 20000,
                            "work_type": "Kirim Barang", "mode": "Air"}, timeout=20)
    assert r.status_code == 200, r.text
    r = requests.post(f"{API}/orders/reconcile",
                      headers=_headers(operation_token),
                      json={"awb_ids": [oid]}, timeout=20)
    assert r.status_code == 200, r.text
    r = requests.post(f"{API}/invoices", headers=_headers(operation_token),
                      json={"awb_id": oid}, timeout=20)
    assert r.status_code in (200, 201), r.text
    inv = r.json()
    return inv["id"], oid


class TestInvoiceExcel:
    @pytest.fixture(scope="class")
    def invoice_bundle(self, operation_token, finance_token, seed_customer):
        return _ensure_invoice(operation_token, finance_token, seed_customer)

    def test_invoice_excel_formatting(self, operation_token, invoice_bundle):
        inv_id, _awb_id = invoice_bundle
        r = requests.get(f"{API}/exports/invoice/{inv_id}",
                         headers=_headers(operation_token), timeout=30)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("Content-Type", ""), r.headers
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active

        # B5 = INVOICE banner
        assert ws["B5"].value == "INVOICE", f"B5={ws['B5'].value}"
        # Font color FFFFFFFF or FFFFFF
        font_color = ws["B5"].font.color
        col_val = font_color.rgb if font_color else None
        assert col_val and str(col_val).upper().endswith("FFFFFF"), f"B5 font color={col_val}"
        # Grey fill 7F7F7F
        fill = ws["B5"].fill
        fg = (fill.fgColor.rgb or "") if fill and fill.fgColor else ""
        assert "7F7F7F" in str(fg).upper(), f"B5 fill fg={fg}"

        # B13 = No (header row), fill 595959
        assert ws["B13"].value == "No", f"B13={ws['B13'].value}"
        fg13 = (ws["B13"].fill.fgColor.rgb or "")
        assert "595959" in str(fg13).upper(), f"B13 fill={fg13}"

        # Indonesian date format check (V8 is invoice date)
        v8 = ws["V8"].value
        assert isinstance(v8, str) and any(m in v8 for m in [
            "Januari", "Februari", "Maret", "April", "Mei", "Juni",
            "Juli", "Agustus", "September", "Oktober", "November", "Desember"
        ]), f"V8 invoice date not Indonesian format: {v8!r}"

        # Summary labels — scan column O (15) for the labels
        labels = []
        for r_idx in range(1, ws.max_row + 1):
            v = ws.cell(r_idx, 15).value
            if isinstance(v, str):
                labels.append(v.strip())
        labels_norm = [l.lstrip(">").lstrip().lstrip(">").strip() for l in labels]
        labels_flat = " | ".join(labels)
        for needed in ["Jumlah", ">> PPN", "Sub Total", ">> Asuransi", ">> Polis", ">> Adm", ">> Materai", "TOTAL"]:
            assert any(needed in lab for lab in labels), f"Missing summary label {needed!r}; got {labels_flat}"

        # BANK ADDRESS label exists somewhere
        all_text = []
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    all_text.append(v)
        text_blob = " | ".join(all_text)
        assert "BANK ADDRESS" in text_blob, f"BANK ADDRESS missing"
        assert "Mandiri" in text_blob, f"default bank_name 'Mandiri' missing: {text_blob[:500]}"


class TestReportArExportFilter:
    def test_report_ar_invoice_ids_filter(self, operation_token, finance_token, seed_customer):
        inv_id, awb_id = _ensure_invoice(operation_token, finance_token, seed_customer)
        # also create a second invoice
        inv2_id, _ = _ensure_invoice(operation_token, finance_token, seed_customer)
        r = requests.get(f"{API}/exports/report-ar?invoice_ids={inv_id}",
                         headers=_headers(operation_token), timeout=30)
        assert r.status_code == 200, r.text
        # Must be xlsx (smoke check)
        assert r.content[:2] == b"PK", "Not a valid xlsx (no PK header)"
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Collect AWB column values (col 5 = AWB based on REPORT_AR_HEADER_ROW1)
        awb_vals = set()
        # Find data rows: row >= 3 typically; iterate and find non-empty col 5
        for row_idx in range(3, ws.max_row + 1):
            v = ws.cell(row_idx, 5).value
            if v:
                awb_vals.add(str(v))
        # Look up our awb_number
        r2 = requests.get(f"{API}/orders/{awb_id}", headers=_headers(operation_token), timeout=20)
        assert r2.status_code == 200
        our_awb = r2.json()["awb_number"]
        assert our_awb in awb_vals, f"filtered AWB {our_awb} not in export; got={awb_vals}"


# ------------------------- Cash advance + breakdown -------------------------
class TestCashAdvanceBreakdown:
    @pytest.fixture(scope="class")
    def three_awbs(self, operation_token, seed_customer):
        oids = []
        for _ in range(3):
            o = _create_order(operation_token, seed_customer)
            oids.append(o["id"])
        return oids

    def _create_advance(self, finance_token, amount=1_000_000) -> dict:
        r = requests.post(f"{API}/cash-advances",
                          headers=_headers(finance_token),
                          json={"recipient_name": "Sopir", "amount": amount,
                                "purpose": "TEST_IT10"}, timeout=20)
        assert r.status_code in (200, 201), r.text
        return r.json()

    def test_settle_with_full_breakdown(self, finance_token, operation_token, three_awbs):
        adv = self._create_advance(finance_token, amount=1_000_000)
        cid = adv["id"]
        breakdown = [
            {"awb_id": three_awbs[0], "amount": 400000},
            {"awb_id": three_awbs[1], "amount": 350000},
            {"awb_id": three_awbs[2], "amount": 250000},
        ]
        r = requests.post(f"{API}/cash-advances/{cid}/settle",
                          headers=_headers(finance_token),
                          json={"actual_amount": 1_000_000, "awb_breakdown": breakdown}, timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("status") == "Settled"
        assert float(body.get("actual_amount") or 0) == 1_000_000
        ab = body.get("awb_breakdown") or []
        assert len(ab) == 3, ab
        for it in ab:
            assert it.get("awb_number"), f"awb_number missing in breakdown item: {it}"

        # Verify costs created
        for it in breakdown:
            rc = requests.get(f"{API}/costs?awb_id={it['awb_id']}",
                              headers=_headers(operation_token), timeout=20)
            costs = rc.json()
            ca_costs = [c for c in costs if c.get("source") == "cash_advance_settle" and c.get("source_ref") == cid]
            assert len(ca_costs) == 1, f"expected 1 cost on AWB {it['awb_id']} from cash advance {cid}, got {ca_costs}"
            assert float(ca_costs[0]["amount"]) == float(it["amount"])
            assert ca_costs[0]["cost_category"] == "Operational"

    def test_settle_breakdown_mismatch_400(self, finance_token, three_awbs):
        adv = self._create_advance(finance_token, amount=500_000)
        cid = adv["id"]
        breakdown = [
            {"awb_id": three_awbs[0], "amount": 100000},
            {"awb_id": three_awbs[1], "amount": 200000},
        ]  # total=300_000 but actual=500_000
        r = requests.post(f"{API}/cash-advances/{cid}/settle",
                          headers=_headers(finance_token),
                          json={"actual_amount": 500_000, "awb_breakdown": breakdown}, timeout=20)
        assert r.status_code == 400, f"expected 400 mismatch, got {r.status_code}: {r.text}"

    def test_settle_with_partial_breakdown_and_refund(self, finance_token, operation_token, three_awbs):
        adv = self._create_advance(finance_token, amount=1_000_000)
        cid = adv["id"]
        breakdown = [
            {"awb_id": three_awbs[0], "amount": 300000},
            {"awb_id": three_awbs[1], "amount": 250000},
            {"awb_id": three_awbs[2], "amount": 150000},
        ]  # total=700_000; refund=300_000; sum=1_000_000
        r = requests.post(f"{API}/cash-advances/{cid}/settle",
                          headers=_headers(finance_token),
                          json={"actual_amount": 700_000, "refund_amount": 300_000,
                                "awb_breakdown": breakdown}, timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("status") == "Settled"
        assert float(body.get("actual_amount") or 0) == 700_000
        assert float(body.get("refund_amount") or 0) == 300_000

        # 3 awb_costs created
        for it in breakdown:
            rc = requests.get(f"{API}/costs?awb_id={it['awb_id']}",
                              headers=_headers(operation_token), timeout=20)
            ca_costs = [c for c in rc.json() if c.get("source") == "cash_advance_settle" and c.get("source_ref") == cid]
            assert len(ca_costs) == 1, f"expected 1 cost for {it['awb_id']}, got {ca_costs}"
            assert float(ca_costs[0]["amount"]) == float(it["amount"])
