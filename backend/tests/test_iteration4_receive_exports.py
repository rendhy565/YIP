"""
Iteration-4 backend tests — Tanggal Terima receive endpoint + Excel exports.

Covers:
  - POST /api/orders/{oid}/receive: full payload persistence + enable/disable flags
  - POST /api/orders/{oid}/receive: invalid date 400, invalid id 404, any-role access
  - GET  /api/exports/report-ar: header layout (row 1 + row 2 sub-headers),
    month filter, group placement of Polis/Insurance/SPSI/mode columns
  - GET  /api/exports/invoice/{invoice_id}: single sheet w/ template cells
  - POST /api/exports/invoices/bulk: multiple sheets, invalid ids skipped, empty 400
  - PUT  /api/settings/signer_PT. Abirama Mitra Pradana -> reflects in invoice signature row
  - GET  /api/settings: seeded signer_/address_/npwp_ keys exist
"""
import io
import os
import uuid
import pytest
import requests
import openpyxl

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"


# ---------- helpers ----------
def _login(u, p):
    r = requests.post(f"{API}/auth/login", json={"username": u, "password": p}, timeout=20)
    assert r.status_code == 200, f"login {u}: {r.status_code} {r.text}"
    return r.json()["access_token"]


def H(t):
    return {"Authorization": f"Bearer {t}", "Content-Type": "application/json"}


@pytest.fixture(scope="session")
def admin_tok(): return _login("admin", "admin123")


@pytest.fixture(scope="session")
def op_tok(): return _login("operation", "operation123")


@pytest.fixture(scope="session")
def fin_tok(): return _login("finance", "finance123")


@pytest.fixture(scope="session")
def first_customer(op_tok):
    r = requests.get(f"{API}/customers", headers=H(op_tok), timeout=20)
    assert r.status_code == 200
    items = r.json()
    if items:
        return items[0]
    rc = requests.post(f"{API}/customers",
                       json={"name": f"TEST_Cust_{uuid.uuid4().hex[:6]}", "address": "Jakarta"},
                       headers=H(op_tok), timeout=20)
    assert rc.status_code == 200, rc.text
    return rc.json()


def _create_order(tok, customer_id, **overrides):
    body = {
        "customer_id": customer_id,
        "item_description": "TEST_Iter4_Item",
        "goods_content": "Box",
        "origin": "JKT",
        "destination": "SBY",
        "tariff_per_kg": 10000,
        "qty_kg": 5,
        "tax_percentage": 11,
        "company_name": "PT. Abirama Mitra Pradana",
        "consignee": "TEST_Penerima",
        "service": "FCL",
        "via": "D",
        "armada": "CDD",
        "nopol": "B 1234 ABC",
        "reff_no": "PO-TEST-001",
        "insurance_enabled": True,
        "insurance_policy": "POL-TEST",
        "insurance_amount": 500_000,
    }
    body.update(overrides)
    r = requests.post(f"{API}/orders", json=body, headers=H(tok), timeout=20)
    assert r.status_code == 200, r.text
    return r.json()


def _open_workbook(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content), data_only=False)


def _create_invoiced_order(op_tok, admin_tok, customer_id, **overrides):
    """Create an order, add a Vendor cost, reconcile, and create an invoice."""
    o = _create_order(op_tok, customer_id, **overrides)
    rc = requests.post(f"{API}/costs", json={
        "awb_id": o["id"],
        "cost_category": "Vendor",
        "amount": 25000,
        "description": "TEST_Iter4_Vendor",
        "vendor_name": "TEST_Vendor",
    }, headers=H(op_tok), timeout=20)
    assert rc.status_code == 200, rc.text
    rr = requests.post(f"{API}/orders/reconcile",
                       json={"awb_ids": [o["id"]]},
                       headers=H(admin_tok), timeout=20)
    assert rr.status_code == 200, rr.text
    ri = requests.post(f"{API}/invoices", json={"awb_id": o["id"]},
                       headers=H(op_tok), timeout=20)
    assert ri.status_code == 200, ri.text
    inv = ri.json()
    return o, inv


# =============================================================================
# 1. Receive endpoint
# =============================================================================
class TestReceiveEndpoint:
    def test_full_payload_persists_any_role(self, op_tok, fin_tok, admin_tok, first_customer):
        # Try with each role to ensure ANY logged user works
        for tok, label in [(op_tok, "op"), (fin_tok, "fin"), (admin_tok, "adm")]:
            o = _create_order(op_tok, first_customer["id"])
            payload = {
                "received_date": "2026-06-15",
                "received_by_name": f"  TEST Receiver {label}  ",
                "polis": "POL-OVERRIDE",
                "asuransi_amount": 750_000,
                "spsi_enabled": True,
                "spsi_amount": 100_000,
                "sewa_forklift_enabled": True,
                "sewa_forklift_amount": 250_000,
                "work_type": "Kirim Barang",
                "mode": "Trucking",
            }
            r = requests.post(f"{API}/orders/{o['id']}/receive",
                              json=payload, headers=H(tok), timeout=20)
            assert r.status_code == 200, f"{label}: {r.status_code} {r.text}"
            d = r.json()
            assert d["received_by_name"] == f"TEST Receiver {label}"
            assert d["received_polis"] == "POL-OVERRIDE"
            assert float(d["received_asuransi_amount"]) == 750_000
            assert d["spsi_enabled"] is True
            assert float(d["spsi_amount"]) == 100_000
            assert d["sewa_forklift_enabled"] is True
            assert float(d["sewa_forklift_amount"]) == 250_000
            assert d["work_type"] == "Kirim Barang"
            assert d["mode"] == "Trucking"
            assert "2026-06-15" in d["received_at"], d["received_at"]

    def test_disabled_flags_zero_amounts(self, op_tok, first_customer):
        o = _create_order(op_tok, first_customer["id"])
        r = requests.post(f"{API}/orders/{o['id']}/receive", json={
            "received_date": "2026-06-20",
            "received_by_name": "TEST_NoExtras",
            "spsi_enabled": False,
            "spsi_amount": 999_999,           # should be ignored
            "sewa_forklift_enabled": False,
            "sewa_forklift_amount": 888_888,  # should be ignored
            "work_type": "Relokasi Barang",
            "mode": "Konsul",
        }, headers=H(op_tok), timeout=20)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["spsi_enabled"] is False
        assert float(d["spsi_amount"]) == 0
        assert d["sewa_forklift_enabled"] is False
        assert float(d["sewa_forklift_amount"]) == 0

    def test_invalid_date_400(self, op_tok, first_customer):
        o = _create_order(op_tok, first_customer["id"])
        r = requests.post(f"{API}/orders/{o['id']}/receive", json={
            "received_date": "not-a-date",
            "received_by_name": "X",
            "work_type": "Kirim Barang",
            "mode": "Air",
        }, headers=H(op_tok), timeout=20)
        assert r.status_code == 400, r.text

    def test_invalid_order_id_404(self, op_tok):
        # Use a syntactically-valid but non-existent ObjectId
        bogus = "0" * 24
        r = requests.post(f"{API}/orders/{bogus}/receive", json={
            "received_by_name": "X",
            "work_type": "Kirim Barang",
            "mode": "Air",
        }, headers=H(op_tok), timeout=20)
        assert r.status_code == 404, r.text


# =============================================================================
# 2. REPORT AR export
# =============================================================================
class TestReportArExport:
    def test_headers_and_groups(self, admin_tok):
        r = requests.get(f"{API}/exports/report-ar", headers=H(admin_tok), timeout=60)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert "REPORT_AR_" in r.headers.get("content-disposition", "")
        assert len(r.content) > 1000

        wb = _open_workbook(r.content)
        assert "REPORT AR" in wb.sheetnames, wb.sheetnames
        ws = wb["REPORT AR"]

        # Row 1 — main cols (1-17)
        expected_r1 = ["Bulan", "Tgl Pick Up", "Tanggal Terima", "Status Pengiriman", "AWB",
                       "Vendor Name", "Transport", "From", "Destination", "No Mobil",
                       "Penerima", "Qty", "Koli", "KG", "Min Kg", "Final KG", "Harga /KG"]
        for i, exp in enumerate(expected_r1, start=1):
            got = ws.cell(1, i).value
            assert got == exp, f"row1 col{i}: expected {exp!r} got {got!r}"

        # Group headers (merged top-left holds the value)
        assert ws.cell(1, 18).value == "Kirim Barang"
        assert ws.cell(1, 27).value == "Dismantle Barang"
        assert ws.cell(1, 36).value == "Relokasi Barang"

        # Trailing single cols 45-50 — accept embedded newlines in 'Amount\n(Exc VAT)'
        amount_cell = (ws.cell(1, 45).value or "").replace("\n", " ").strip()
        assert "Amount" in amount_cell and "Exc VAT" in amount_cell, amount_cell
        assert ws.cell(1, 46).value == "Keterangan"
        assert ws.cell(1, 47).value == "Harga"
        assert ws.cell(1, 48).value == "List Muatan"
        assert ws.cell(1, 49).value == "Asuransi"
        assert ws.cell(1, 50).value == "No Invoice"

        # Row 2 — sub-headers for the 3 groups
        kirim_subs = ["Polis", "Insurance", "SPSI", "Sewa Froklift", "Kirim Barang",
                      "Loading/Unloading", "Air", "Konsul", "Trucking"]
        for off, exp in enumerate(kirim_subs):
            assert ws.cell(2, 18 + off).value == exp, \
                f"kirim sub col {18+off}: {ws.cell(2, 18+off).value!r} != {exp!r}"
        dis_subs = ["Polis", "Insurance", "SPSI", "Sewa Froklift", "Packing",
                    "Loading/Unloading", "Air", "Konsul", "Trucking"]
        for off, exp in enumerate(dis_subs):
            assert ws.cell(2, 27 + off).value == exp
        rel_subs = ["Polis", "Insurance", "SPSI", "Sewa Froklift", "Loading/Unloading",
                    "Packing", "Air", "Konsul", "Trucking"]
        for off, exp in enumerate(rel_subs):
            assert ws.cell(2, 36 + off).value == exp

    def test_month_filter(self, admin_tok):
        r_all = requests.get(f"{API}/exports/report-ar", headers=H(admin_tok), timeout=60)
        r_jun = requests.get(f"{API}/exports/report-ar?month=2026-06",
                             headers=H(admin_tok), timeout=60)
        assert r_all.status_code == 200
        assert r_jun.status_code == 200
        wb_all = _open_workbook(r_all.content)["REPORT AR"]
        wb_jun = _open_workbook(r_jun.content)["REPORT AR"]
        rows_all = sum(1 for row in wb_all.iter_rows(min_row=3, max_col=5) if row[4].value)
        rows_jun = sum(1 for row in wb_jun.iter_rows(min_row=3, max_col=5) if row[4].value)
        # June filter should not produce MORE rows than all
        assert rows_jun <= rows_all
        # All non-empty June rows should have "JUN" in col 1 (Bulan)
        for row in wb_jun.iter_rows(min_row=3, max_col=5):
            if row[4].value:
                bulan = row[0].value
                assert bulan == "JUN", f"month-filter leaked: bulan={bulan!r}, awb={row[4].value!r}"

    def test_invalid_month_400(self, admin_tok):
        r = requests.get(f"{API}/exports/report-ar?month=bad",
                         headers=H(admin_tok), timeout=20)
        assert r.status_code == 400

    def test_group_placement_kirim_trucking(self, admin_tok, op_tok, first_customer):
        # Create + receive an AWB with Kirim/Trucking, then verify columns
        o = _create_order(op_tok, first_customer["id"],
                          tariff_per_kg=20_000, qty_kg=3)  # grand_total = 60000 + 11% = 66600
        polis_val = f"POL-K-{uuid.uuid4().hex[:6]}"
        rr = requests.post(f"{API}/orders/{o['id']}/receive", json={
            "received_date": "2026-06-10",
            "received_by_name": "TEST_KirimTruck",
            "polis": polis_val,
            "asuransi_amount": 333_000,
            "spsi_enabled": True,
            "spsi_amount": 123_000,
            "sewa_forklift_enabled": False,
            "work_type": "Kirim Barang",
            "mode": "Trucking",
        }, headers=H(op_tok), timeout=20)
        assert rr.status_code == 200, rr.text
        grand_total = float(o["grand_total"])

        r = requests.get(f"{API}/exports/report-ar", headers=H(admin_tok), timeout=60)
        assert r.status_code == 200
        ws = _open_workbook(r.content)["REPORT AR"]
        target_row = None
        for row in ws.iter_rows(min_row=3, max_col=5):
            if row[4].value == o["awb_number"]:
                target_row = row[0].row
                break
        assert target_row, f"AWB {o['awb_number']} not found in REPORT AR"
        assert ws.cell(target_row, 18).value == polis_val
        assert float(ws.cell(target_row, 19).value) == 333_000  # Insurance
        assert float(ws.cell(target_row, 20).value) == 123_000  # SPSI
        # Sewa Froklift disabled → blank
        assert ws.cell(target_row, 21).value in (None, "")
        # Mode = Trucking → offset 8 → col 18+8 = 26
        assert float(ws.cell(target_row, 26).value) == grand_total
        # Other mode cols (Air=24, Konsul=25) must be empty
        for col in (24, 25):
            assert ws.cell(target_row, col).value in (None, ""), \
                f"col {col} should be empty for Trucking, got {ws.cell(target_row, col).value!r}"

    def test_group_placement_dismantle_air(self, admin_tok, op_tok, first_customer):
        o = _create_order(op_tok, first_customer["id"], tariff_per_kg=15_000, qty_kg=4)
        polis_val = f"POL-D-{uuid.uuid4().hex[:6]}"
        rr = requests.post(f"{API}/orders/{o['id']}/receive", json={
            "received_date": "2026-06-11",
            "received_by_name": "TEST_DismAir",
            "polis": polis_val,
            "asuransi_amount": 444_000,
            "spsi_enabled": False,
            "sewa_forklift_enabled": False,
            "work_type": "Dismantle Barang",
            "mode": "Air",
        }, headers=H(op_tok), timeout=20)
        assert rr.status_code == 200, rr.text
        grand_total = float(o["grand_total"])

        r = requests.get(f"{API}/exports/report-ar", headers=H(admin_tok), timeout=60)
        ws = _open_workbook(r.content)["REPORT AR"]
        target_row = None
        for row in ws.iter_rows(min_row=3, max_col=5):
            if row[4].value == o["awb_number"]:
                target_row = row[0].row
                break
        assert target_row, "AWB row not found"
        # Dismantle group: gstart=27 → Polis=27, Insurance=28, Air=33 (offset 6)
        assert ws.cell(target_row, 27).value == polis_val
        assert float(ws.cell(target_row, 28).value) == 444_000
        assert float(ws.cell(target_row, 33).value) == grand_total
        # Kirim/Relokasi groups should be empty for this row
        for col in (18, 19, 20, 26, 36, 42, 43, 44):
            v = ws.cell(target_row, col).value
            assert v in (None, ""), f"col {col} should be empty for dismantle row, got {v!r}"


# =============================================================================
# 3. Invoice export (single + bulk)
# =============================================================================
class TestInvoiceExport:
    @pytest.fixture(scope="class")
    def invoiced(self, op_tok, admin_tok, first_customer):
        o, inv = _create_invoiced_order(op_tok, admin_tok, first_customer["id"],
                                        company_name="PT. Abirama Mitra Pradana")
        return o, inv

    def test_single_invoice_export(self, admin_tok, invoiced):
        o, inv = invoiced
        r = requests.get(f"{API}/exports/invoice/{inv['id']}",
                         headers=H(admin_tok), timeout=30)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert "Invoice_" in r.headers.get("content-disposition", "")
        wb = _open_workbook(r.content)
        # Sheet title = last 31 chars of invoice number
        expected_title = (inv["invoice_number"])[-31:]
        assert expected_title in wb.sheetnames, wb.sheetnames
        ws = wb[expected_title]
        assert ws["B6"].value == "INVOICE"
        assert ws["B8"].value == "PT. Abirama Mitra Pradana"
        # R8 = invoice_number
        assert ws["R8"].value == inv["invoice_number"]
        # Attn = consignee
        assert ws["C13"].value == o["consignee"]

    def test_bulk_invoice_export_multiple_sheets(self, op_tok, admin_tok, first_customer):
        # Create 2 invoiced orders
        _, inv_a = _create_invoiced_order(op_tok, admin_tok, first_customer["id"])
        _, inv_b = _create_invoiced_order(op_tok, admin_tok, first_customer["id"])
        r = requests.post(f"{API}/exports/invoices/bulk",
                          json={"invoice_ids": [inv_a["id"], inv_b["id"]]},
                          headers=H(admin_tok), timeout=30)
        assert r.status_code == 200, r.text
        wb = _open_workbook(r.content)
        assert len(wb.sheetnames) >= 2
        for inv in (inv_a, inv_b):
            assert any(inv["invoice_number"].endswith(s) or s == inv["invoice_number"][-31:]
                       for s in wb.sheetnames), wb.sheetnames

    def test_bulk_invalid_ids_skipped(self, op_tok, admin_tok, first_customer):
        _, inv_a = _create_invoiced_order(op_tok, admin_tok, first_customer["id"])
        bogus = "0" * 24
        r = requests.post(f"{API}/exports/invoices/bulk",
                          json={"invoice_ids": [inv_a["id"], bogus, "not-an-objectid"]},
                          headers=H(admin_tok), timeout=30)
        assert r.status_code == 200, r.text
        wb = _open_workbook(r.content)
        # Only the valid one should produce a sheet
        assert len(wb.sheetnames) == 1, wb.sheetnames

    def test_bulk_empty_400(self, admin_tok):
        r = requests.post(f"{API}/exports/invoices/bulk",
                          json={"invoice_ids": []},
                          headers=H(admin_tok), timeout=20)
        assert r.status_code == 400, r.text


# =============================================================================
# 4. Settings: seeded keys + signer PUT reflected in invoice
# =============================================================================
class TestSettingsAndSigner:
    EXPECTED_KEYS = [
        "signer_PT. Abirama Mitra Pradana",
        "signer_PT. Abirama Mitra Pratama",
        "address_PT. Abirama Mitra Pradana",
        "address_PT. Abirama Mitra Pratama",
        "npwp_PT. Abirama Mitra Pradana",
        "npwp_PT. Abirama Mitra Pratama",
    ]

    def test_seeded_keys_present(self, admin_tok):
        r = requests.get(f"{API}/settings", headers=H(admin_tok), timeout=20)
        assert r.status_code == 200, r.text
        keys = {s["key_name"] for s in r.json()}
        missing = [k for k in self.EXPECTED_KEYS if k not in keys]
        assert not missing, f"missing seed settings: {missing}"

    def test_signer_put_reflects_in_invoice(self, op_tok, admin_tok, first_customer):
        # 1) Set signer
        signer_value = f"TEST JOHN DOE {uuid.uuid4().hex[:4].upper()}"
        key = "signer_PT. Abirama Mitra Pradana"
        r_put = requests.put(f"{API}/settings/{key}",
                             json={"key_name": key, "key_value": signer_value},
                             headers=H(admin_tok), timeout=20)
        assert r_put.status_code == 200, r_put.text

        # Verify GET reflects
        r_get = requests.get(f"{API}/settings", headers=H(admin_tok), timeout=20)
        got = {s["key_name"]: s.get("key_value") for s in r_get.json()}
        assert got.get(key) == signer_value

        # 2) Create invoice for Pradana company AWB
        o, inv = _create_invoiced_order(op_tok, admin_tok, first_customer["id"],
                                        company_name="PT. Abirama Mitra Pradana")
        # 3) Export and scan worksheet for signer text
        r_exp = requests.get(f"{API}/exports/invoice/{inv['id']}",
                             headers=H(admin_tok), timeout=30)
        assert r_exp.status_code == 200, r_exp.text
        wb = _open_workbook(r_exp.content)
        ws = wb[wb.sheetnames[0]]
        # Scan all cells for the signer name
        found = False
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and signer_value in v:
                    found = True
                    break
            if found:
                break
        assert found, f"signer {signer_value!r} not found in invoice workbook"
