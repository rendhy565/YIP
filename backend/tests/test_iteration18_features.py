"""
Iteration 18 — Tests covering:
- ReceiveAwbIn.received_notes (catatan) — optional + saved to awb_orders.received_notes
- PATCH /api/invoices/{iid}/po-number (set + clear) → updates po_number + po_number_updated_at
- GET /api/invoices/po-number-template → valid xlsx with proper headers
- POST /api/invoices/import-po-numbers → updates matched, returns not_found list, idempotent
- GET /api/invoices/{iid}/full → awbs[] (multi-AWB safe) + invoice.po_number present
- GET /api/invoices → list contains po_number field per invoice
"""
import io
import os
import uuid
import openpyxl
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")

def _login(u, p):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": u, "password": p}, timeout=30)
    assert r.status_code == 200, f"Login {u} failed: {r.text}"
    return r.json()["access_token"]

def H(tok): return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}
def HF(tok): return {"Authorization": f"Bearer {tok}"}  # for multipart

@pytest.fixture(scope="module")
def admin_token(): return _login("admin", "admin123")
@pytest.fixture(scope="module")
def op_token(): return _login("operation", "operation123")

@pytest.fixture(scope="module")
def cust(admin_token):
    r = requests.get(f"{BASE_URL}/api/customers", headers=H(admin_token), timeout=30)
    cs = r.json() if r.status_code == 200 else []
    if cs: return cs[0]
    cr = requests.post(f"{BASE_URL}/api/customers", headers=H(admin_token),
                       json={"name": f"TEST_IT18_C_{uuid.uuid4().hex[:5]}", "address": "Jl."}, timeout=30)
    return cr.json()


def _mk_awb(tok, cust, tariff=5000, qty=100):
    payload = {
        "customer_id": cust["id"],
        "company_name": cust.get("default_company") or "PT. Abirama Mitra Pradana",
        "item_description": "TEST-IT18", "goods_content": "Box",
        "origin": "JKT", "destination": "SBY",
        "tariff_per_kg": tariff, "qty_kg": qty, "tax_percentage": 0,
        "service": "LCL", "via": "D", "koli_count": 1,
    }
    r = requests.post(f"{BASE_URL}/api/orders", headers=H(tok), json=payload, timeout=30)
    assert r.status_code in (200, 201), r.text
    return r.json()

def _today():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")

def _add_cost(tok, awb_id, amount=10000):
    p = {"awb_id": awb_id, "cost_category": "Vendor", "sub_label": "Trucking",
         "vendor_name": f"TEST_V_{uuid.uuid4().hex[:4]}", "amount": amount, "description": "x"}
    r = requests.post(f"{BASE_URL}/api/costs", headers=H(tok), json=p, timeout=30)
    assert r.status_code in (200, 201), r.text
    return r.json()

def _full_received_and_pod(tok, awb_id, with_notes=True):
    body = {
        "received_by_name": "Tester",
        "received_date": _today(),
        "spsi_enabled": False, "sewa_forklift_enabled": False,
        "polis_enabled": False, "asuransi_enabled": False,
        "work_type": "Kirim Barang", "mode": "Trucking",
    }
    if with_notes:
        body["received_notes"] = "catatan test"
    r = requests.post(f"{BASE_URL}/api/orders/{awb_id}/receive", headers=H(tok), json=body, timeout=30)
    assert r.status_code == 200, r.text
    rp = requests.post(f"{BASE_URL}/api/orders/{awb_id}/receive-pod", headers=H(tok),
                       json={"pod_date": _today(), "pod_signed_by": "POD-Signer", "pod_notes": "ok"}, timeout=30)
    assert rp.status_code == 200, rp.text
    return r.json()


# =====================================================================
# Test 1: received_notes (catatan) — optional + saved
# =====================================================================
class TestReceiveNotes:
    def test_receive_with_notes_saves(self, op_token, cust):
        awb = _mk_awb(op_token, cust)
        r = requests.post(f"{BASE_URL}/api/orders/{awb['id']}/receive", headers=H(op_token), json={
            "received_by_name": "X",
            "spsi_enabled": False, "sewa_forklift_enabled": False,
            "polis_enabled": False, "asuransi_enabled": False,
            "received_notes": "catatan test",
            "work_type": "Kirim Barang", "mode": "Trucking",
        }, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data.get("received_notes") == "catatan test", f"received_notes not saved: {data.get('received_notes')}"

    def test_receive_without_notes_optional(self, op_token, cust):
        awb = _mk_awb(op_token, cust)
        r = requests.post(f"{BASE_URL}/api/orders/{awb['id']}/receive", headers=H(op_token), json={
            "received_by_name": "X",
            "spsi_enabled": False, "sewa_forklift_enabled": False,
            "polis_enabled": False, "asuransi_enabled": False,
            "work_type": "Kirim Barang", "mode": "Trucking",
        }, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        # received_notes must be None or absent (not a 422 error)
        assert data.get("received_notes") in (None, "", ), f"received_notes should be None, got: {data.get('received_notes')!r}"


# =====================================================================
# Helper: create a multi-AWB invoice via grouped bulk
# =====================================================================
@pytest.fixture(scope="module")
def multi_awb_invoice(op_token, cust):
    """Create 2 AWBs from same customer → reconcile both → bulk invoice grouped."""
    awb1 = _mk_awb(op_token, cust)
    awb2 = _mk_awb(op_token, cust)
    _add_cost(op_token, awb1["id"])
    _add_cost(op_token, awb2["id"])
    _full_received_and_pod(op_token, awb1["id"])
    _full_received_and_pod(op_token, awb2["id"])
    rec = requests.post(f"{BASE_URL}/api/orders/reconcile", headers=H(op_token),
                        json={"awb_ids": [awb1["id"], awb2["id"]]}, timeout=30)
    assert rec.status_code == 200, rec.text
    # Bulk invoice with grouped
    rb = requests.post(f"{BASE_URL}/api/invoices/bulk", headers=H(op_token),
                       json={"awb_ids": [awb1["id"], awb2["id"]],
                             "grouped": [[awb1["id"], awb2["id"]]]}, timeout=30)
    assert rb.status_code == 200, rb.text
    created = rb.json().get("created") or []
    assert len(created) >= 1, f"No invoice created: {rb.text}"
    inv_id = created[0]["id"]
    inv_number = created[0]["invoice_number"]
    return {"id": inv_id, "invoice_number": inv_number, "awb_ids": [awb1["id"], awb2["id"]]}


# =====================================================================
# Test 2: PATCH /api/invoices/{iid}/po-number (set + clear)
# =====================================================================
class TestPOPatch:
    def test_set_po_number(self, op_token, multi_awb_invoice):
        iid = multi_awb_invoice["id"]
        r = requests.patch(f"{BASE_URL}/api/invoices/{iid}/po-number", headers=H(op_token),
                           json={"po_number": "PO-123"}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("po_number") == "PO-123"
        assert d.get("po_number_updated_at") is not None
        # GET verify persisted
        gf = requests.get(f"{BASE_URL}/api/invoices/{iid}/full", headers=H(op_token), timeout=30)
        assert gf.status_code == 200
        assert gf.json()["invoice"].get("po_number") == "PO-123"

    def test_clear_po_number(self, op_token, multi_awb_invoice):
        iid = multi_awb_invoice["id"]
        # First set
        requests.patch(f"{BASE_URL}/api/invoices/{iid}/po-number", headers=H(op_token),
                       json={"po_number": "PO-TO-CLEAR"}, timeout=30)
        # Then clear
        r = requests.patch(f"{BASE_URL}/api/invoices/{iid}/po-number", headers=H(op_token),
                           json={"po_number": None}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("po_number") in (None, ""), f"expected None, got: {d.get('po_number')!r}"


# =====================================================================
# Test 3: GET /api/invoices/po-number-template → valid xlsx
# =====================================================================
class TestPOTemplate:
    def test_template_returns_xlsx(self, op_token):
        r = requests.get(f"{BASE_URL}/api/invoices/po-number-template", headers=H(op_token), timeout=30)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheet" in ct or "officedocument" in ct, f"unexpected content-type: {ct}"
        # Parse with openpyxl to confirm valid file
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
        ws = wb.active
        h1 = (ws.cell(1, 1).value or "").strip()
        h2 = (ws.cell(1, 2).value or "").strip()
        assert "Invoice Number" in h1, f"Header A1 missing 'Invoice Number': {h1!r}"
        assert "PO Number" in h2, f"Header B1 missing 'PO Number': {h2!r}"


# =====================================================================
# Test 4: POST /api/invoices/import-po-numbers — matched + not_found + idempotency
# =====================================================================
class TestPOImport:
    def _make_xlsx(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(1, 1, "Invoice Number *")
        ws.cell(1, 2, "PO Number")
        for i, (inv_no, po_no) in enumerate(rows, start=2):
            ws.cell(i, 1, inv_no)
            ws.cell(i, 2, po_no)
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio

    def test_import_matched_and_not_found(self, op_token, multi_awb_invoice):
        inv_no = multi_awb_invoice["invoice_number"]
        fake_no = f"INVOICE-DOES-NOT-EXIST-{uuid.uuid4().hex[:6]}"
        f = self._make_xlsx([(inv_no, "PO-IMPORT-XYZ"), (fake_no, "PO-ORPHAN")])
        r = requests.post(f"{BASE_URL}/api/invoices/import-po-numbers", headers=HF(op_token),
                          files={"file": ("po.xlsx", f.getvalue(),
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                          timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d.get("ok") is True
        assert d.get("updated_count") == 1, f"expected 1 updated, got {d.get('updated_count')}: {d}"
        assert d.get("not_found_count") == 1, f"expected 1 not_found, got {d.get('not_found_count')}: {d}"
        assert fake_no in (d.get("not_found") or [])
        # Verify persisted on full endpoint
        gf = requests.get(f"{BASE_URL}/api/invoices/{multi_awb_invoice['id']}/full",
                          headers=H(op_token), timeout=30)
        assert gf.json()["invoice"].get("po_number") == "PO-IMPORT-XYZ"

    def test_import_idempotent(self, op_token, multi_awb_invoice):
        inv_no = multi_awb_invoice["invoice_number"]
        f1 = self._make_xlsx([(inv_no, "PO-IDEMPOTENT-A")])
        r1 = requests.post(f"{BASE_URL}/api/invoices/import-po-numbers", headers=HF(op_token),
                           files={"file": ("po.xlsx", f1.getvalue(),
                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                           timeout=30)
        assert r1.status_code == 200, r1.text
        d1 = r1.json()
        # second identical upload
        f2 = self._make_xlsx([(inv_no, "PO-IDEMPOTENT-A")])
        r2 = requests.post(f"{BASE_URL}/api/invoices/import-po-numbers", headers=HF(op_token),
                           files={"file": ("po.xlsx", f2.getvalue(),
                                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                           timeout=30)
        assert r2.status_code == 200, r2.text
        d2 = r2.json()
        assert d1.get("updated_count") == d2.get("updated_count") == 1
        assert d1.get("not_found_count") == d2.get("not_found_count") == 0


# =====================================================================
# Test 5: GET /api/invoices/{iid}/full → awbs[] array + invoice.po_number
# =====================================================================
class TestInvoiceFullMultiAWB:
    def test_full_returns_awbs_array(self, op_token, multi_awb_invoice):
        iid = multi_awb_invoice["id"]
        # Ensure po_number set
        requests.patch(f"{BASE_URL}/api/invoices/{iid}/po-number", headers=H(op_token),
                       json={"po_number": "PO-FULL-CHK"}, timeout=30)
        r = requests.get(f"{BASE_URL}/api/invoices/{iid}/full", headers=H(op_token), timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert "awbs" in d, f"awbs missing: keys={list(d.keys())}"
        assert isinstance(d["awbs"], list)
        assert len(d["awbs"]) == 2, f"expected 2 awbs, got {len(d['awbs'])}"
        # invoice.po_number present
        assert d["invoice"].get("po_number") == "PO-FULL-CHK"
        # Each awb has awb_number
        for a in d["awbs"]:
            assert a.get("awb_number"), f"awb missing awb_number: {a}"


# =====================================================================
# Test 6: GET /api/invoices list — po_number field present per invoice
# =====================================================================
class TestInvoiceListPONumber:
    def test_list_invoices_contains_po_number(self, op_token, multi_awb_invoice):
        iid = multi_awb_invoice["id"]
        # set explicit po
        requests.patch(f"{BASE_URL}/api/invoices/{iid}/po-number", headers=H(op_token),
                       json={"po_number": "PO-LIST-OK"}, timeout=30)
        r = requests.get(f"{BASE_URL}/api/invoices", headers=H(op_token), timeout=30)
        assert r.status_code == 200, r.text
        lst = r.json()
        assert isinstance(lst, list)
        match = next((x for x in lst if x.get("id") == iid), None)
        assert match is not None, f"invoice {iid} not in list"
        assert match.get("po_number") == "PO-LIST-OK", f"po_number mismatch: {match.get('po_number')!r}"
